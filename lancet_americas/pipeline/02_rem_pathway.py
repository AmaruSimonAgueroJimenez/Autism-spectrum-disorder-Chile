#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""02_rem_pathway.py — ruta administrativa agregada REM (A03, A27, A05, A28, P2, P6), Chile 2019–2025.

Construye, a partir de REM Serie A (flujo mensual) y Serie P (stock semestral), la tabla tidy
establecimiento × mes × código con los estados «valor», «cero» y «vacío» diferenciados (la ausencia de fila
es «no reportado» y nunca se imputa), los agregados nacionales anuales por código y por variante de definición
(`config.VARIANTS`: con_rett / sin_rett, más autismo estricto y TGD amplio pre-2021), la tabla edad × sexo de
A05, el panel establecimiento-año y la verificación de cada código contra el diccionario anual oficial.

Reglas (ver PROMPT_CODEX_LANCET_AMERICAS.md y scripts/downloads/DATA_REVIEW.md):
  * Serie A es flujo: total anual = suma de meses. A03 usa COL01 (hombres) + COL02 (mujeres) sumando las celdas
    presentes; A05, A27 y A28 usan COL01 (total / nº de intervenciones).
  * Serie P es stock: diciembre (MES=12) es el análisis principal y junio (MES=06) una sensibilidad separada;
    nunca se suman semestres.
  * Las filas se conservan tal como vienen en el archivo (sin drop_duplicates): reproduce los totales de control.
    Los duplicados exactos se marcan y se informa un total alternativo sin ellos (`total_dedup`).
  * Los conteos son reconocimiento administrativo, no prevalencia ni incidencia; las fuentes no se enlazan por persona.

Ejecución (desde la raíz del repositorio):
    python3 lancet_americas/pipeline/02_rem_pathway.py [--years 2019 2025] [--chunksize N]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import config as CFG  # noqa: E402
import common as C  # noqa: E402

MODULE = "02_rem_pathway"
SCRIPT = "lancet_americas/pipeline/02_rem_pathway.py"
CONTROLS_DIR = CFG.OUT / "controls"  # config.CONTROLS es el diccionario de valores esperados; el directorio es outputs/controls
CONTROLS_DIR.mkdir(parents=True, exist_ok=True)
RAW_COLS = [f"Col{i:02d}" for i in range(1, 51)]
TIDY_COLS = [c.lower() for c in RAW_COLS]
KEYS = ["Ano", "Mes", "IdServicio", "IdEstablecimiento", "CodigoPrestacion", "IdRegion", "IdComuna"]
AGE_GROUPS = C.AGE_GROUPS  # 17 grupos OMS: 0-4 … 75-79, 80+
SEX_LABELS = {"H": "Hombres", "M": "Mujeres"}

# Controles adicionales tomados de scripts/downloads/DATA_REVIEW.md (no están en config.CONTROLS).
EXTRA_CONTROLS = {
    "p2_tea_june_2020": {"code": "P2500500", "measure": "june_stock", "field": "total", "values": {2020: 172}},
    "p2_establishments_june_2020": {"code": "P2500500", "measure": "june_stock", "field": "n_reporting_establishments", "values": {2020: 28}},
    "a03_2023_2024_high_referred": {"code": "09600216", "measure": "annual_sum", "field": "total", "values": {2023: 1152, 2024: 1741}},
    "a03_2023_2024_second_no_referral": {"code": "09600218", "measure": "annual_sum", "field": "total", "values": {2023: 817, 2024: 2341}},
    "a03_2023_2024_second_referral": {"code": "09600219", "measure": "annual_sum", "field": "total", "values": {2023: 1136, 2024: 2283}},
    "a03_2025_low": {"code": "03710016", "measure": "annual_sum", "field": "total", "values": {2025: 12971}},
    "a03_2025_medium_no_referral": {"code": "03710017", "measure": "annual_sum", "field": "total", "values": {2025: 1637}},
    "a03_2025_medium_referral": {"code": "03710018", "measure": "annual_sum", "field": "total", "values": {2025: 2500}},
    "a03_2025_high_referral": {"code": "03710019", "measure": "annual_sum", "field": "total", "values": {2025: 2181}},
    "a03_2025_susp_30_59_no_referral": {"code": "03710020", "measure": "annual_sum", "field": "total", "values": {2025: 5140}},
    "a03_2025_susp_30_59_referral": {"code": "03710021", "measure": "annual_sum", "field": "total", "values": {2025: 5883}},
}


def log(msg: str) -> None:
    print(f"[{MODULE}] {msg}", flush=True)


# ---------------------------------------------------------------------------
# 1. Catálogo de códigos (CSV verificado + conjuntos de config)
# ---------------------------------------------------------------------------
def load_catalogue() -> pd.DataFrame:
    cat = pd.read_csv(CFG.PATHS["rem_pathway_codes"], dtype=str)
    cat["year_start"] = cat.year_start.astype(int)
    cat["year_end"] = cat.year_end.astype(int)
    cat["series"] = np.where(cat.module.str.startswith("P"), "P", "A")
    cat["era"] = cat.year_start.astype(str) + "–" + cat.year_end.astype(str)
    config_codes = {
        "A05": set(CFG.A05_ENTRY.values()) | set(CFG.A05_EXIT.values()) | set(CFG.A05_BROAD_PRE2021.values()),
        "A27": set(CFG.A27.values()), "A28": set(CFG.A28.values()),
        "A03": set(CFG.A03_LEGACY.values()) | set(CFG.A03_2023_2024.values()) | set(CFG.A03_2024_31_59.values()) | set(CFG.A03_2025.values()),
        "P2": {CFG.P2_TEA, CFG.P2_NANEAS_TOTAL},
        "P6": set(CFG.P6_PRIMARY.values()) | set(CFG.P6_SPECIALTY.values()) | set(CFG.P6_BROAD_PRE2021.values()),
    }
    missing = [(m, c) for m, cs in config_codes.items() for c in cs if c not in set(cat.loc[cat.module == m, "code"])]
    if missing:
        raise ValueError(f"Códigos de config.py ausentes en rem_pathway_codes.csv: {missing}")
    if cat.duplicated(["module", "code"]).any():
        raise ValueError("rem_pathway_codes.csv contiene códigos repetidos dentro de un módulo")
    return cat


# ---------------------------------------------------------------------------
# 2. Verificación contra los diccionarios anuales
# ---------------------------------------------------------------------------
def dictionary_path(series: str, year: int) -> Path | None:
    if series == "A":
        d = CFG.PATHS["rem_a_dicts"] / str(year)
        pat = ("SA", "DICCIONARIO CODIGOS SA")
    else:
        d = CFG.PATHS["rem_p_dicts"] / str(year)
        pat = ("SP", "DICCIONARIO CODIGOS SP")
    if not d.is_dir():
        return None
    files = [f for f in sorted(os.listdir(d)) if f.startswith(pat) and f.lower().endswith((".xlsm", ".xlsx", ".xls")) and not f.startswith("~")]
    if len(files) != 1:
        raise ValueError(f"Diccionario Serie {series} {year} no unívoco: {files}")
    return d / files[0]


def _ascii(s: str) -> str:
    import unicodedata
    return "".join(ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch)).lower()


def _clean(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v).replace("\n", " ")).strip()
    return s or None


def _is_text(v) -> bool:
    s = _clean(v)
    if s is None or s.startswith("#"):
        return False
    try:
        float(s.replace(",", "."))
        return False
    except ValueError:
        return True


def _norm_code(v, series: str) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if series == "A" and re.fullmatch(r"\d+", s):
        s = s.zfill(8)
    return s


def _is_data_row(row, series: str) -> bool:
    c = _norm_code(row[0], series)
    ok = bool(re.fullmatch(r"\d{8}", c)) if series == "A" else bool(re.fullmatch(r"P\d{7}", c))
    return ok and any(_clean(v) == "COL01" for v in row[1:])


_SHEET_CACHE: dict = {}


def load_sheet(path: Path, sheet: str, ncols: int = 60):
    key = (str(path), sheet)
    if key not in _SHEET_CACHE:
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = [tuple(r) for r in wb[sheet].iter_rows(max_col=ncols, values_only=True)] if sheet in wb.sheetnames else None
            wb.close()
        _SHEET_CACHE[key] = rows
    return _SHEET_CACHE[key]


def header_matrix(rows, start: int, end: int, series: str) -> list[list]:
    """Filas de encabezado entre la sección y la fila del código, con relleno jerárquico hacia la derecha
    (una etiqueta se propaga a celdas vacías solo mientras los encabezados superiores no cambien de grupo)."""
    hdr: list[list] = []
    for j in range(start, end):
        r = rows[j]
        if _is_data_row(r, series):
            continue
        area = [_clean(v) if _is_text(v) else None for v in r[2:]]
        if not any(area):
            continue
        ff = [None] * len(area)
        last, src = None, None
        for c, v in enumerate(area):
            if v is not None:
                last, src = v, c
                ff[c] = v
            elif last is not None and all(h[c] == h[src] for h in hdr):
                ff[c] = last
            else:
                last, src = None, None
        hdr.append(ff)
    return hdr


def col_meaning(hdr, pos: int) -> str:
    idx = pos - 2
    return " | ".join(h[idx] for h in hdr if 0 <= idx < len(h) and h[idx])


def check_code(series: str, year: int, sheet: str, code: str, in_era: bool) -> dict:
    rec = dict(year=year, series=series, code=code, sheet=sheet, in_era=in_era, dictionary_file="", excel_row=np.nan, label="", section="",
               col01_meaning="", col02_meaning="", col03_meaning="", col04_meaning="", col37_meaning="", n_col_tokens=np.nan,
               age_columns="", layout_check="", found_in_dictionary=False)
    path = dictionary_path(series, year)
    if path is None:
        rec["layout_check"] = "dictionary folder missing"
        return rec
    rec["dictionary_file"] = str(path.relative_to(CFG.DATA_ROOT))
    rows = load_sheet(path, sheet)
    if rows is None:
        rec["layout_check"] = f"sheet {sheet} missing"
        return rec
    idx = next((i for i, r in enumerate(rows) if _norm_code(r[0], series) == code and _is_data_row(r, series)), None)
    if idx is None:
        return rec
    row = rows[idx]
    sec = next((j for j in range(idx, -1, -1) if "SECCI" in " | ".join(_clean(v) or "" for v in rows[j][:4]).upper()), None)
    tokens = [_clean(v) for v in row]
    first_col = next(k for k, v in enumerate(tokens) if v == "COL01")
    hdr = header_matrix(rows, (sec + 1) if sec is not None else max(0, idx - 8), idx, series)
    positions = {v: k for k, v in enumerate(tokens) if v and re.fullmatch(r"COL\d\d", v)}
    meaning = {c: (col_meaning(hdr, p) if c in positions else "absent") for c, p in [(c, positions.get(c)) for c in ("COL01", "COL02", "COL03", "COL04", "COL37")]}
    label = " | ".join(t for t in tokens[1:first_col] if t and not t.startswith("COL"))
    if tokens[1] is None:  # etiqueta de grupo en celda combinada de una fila anterior (p. ej. «Sí»/«No», «Riesgo medio»)
        j = idx - 1
        while j > (sec if sec is not None else -1) and _clean(rows[j][1]) is None:
            j -= 1
        if j > (sec if sec is not None else -1) and _is_data_row(rows[j], series):
            label = f"[{_clean(rows[j][1])}] {label}"
    rec.update(found_in_dictionary=True, excel_row=idx + 1, label=label,
               section=" | ".join(_clean(v) for v in rows[sec][:4] if _clean(v)) if sec is not None else "",
               n_col_tokens=len(positions), col01_meaning=meaning["COL01"], col02_meaning=meaning["COL02"], col03_meaning=meaning["COL03"],
               col04_meaning=meaning["COL04"], col37_meaning=meaning["COL37"])
    # Verificación de la disposición de columnas según el módulo.
    problems = []
    if sheet == "A03":
        if "hombres" not in meaning["COL01"].lower() or "mujeres" not in meaning["COL02"].lower():
            problems.append("COL01/COL02 no son Hombres/Mujeres")
    elif sheet == "A27":
        if "intervencion" not in _ascii(meaning["COL01"]):
            problems.append("COL01 no es Nº de intervenciones")
    else:  # A05, A28, P2, P6: COL01 total, COL02/COL03 sexo, COL04..COL37 edad × sexo
        if "ambos" not in meaning["COL01"].lower():
            problems.append("COL01 no es Ambos sexos")
        if "hombres" not in meaning["COL02"].lower() or "mujeres" not in meaning["COL03"].lower():
            problems.append("COL02/COL03 no son Hombres/Mujeres")
        ages = []
        for k in range(17):
            ch, cm = f"COL{4 + 2 * k:02d}", f"COL{5 + 2 * k:02d}"
            mh = col_meaning(hdr, positions[ch]).lower() if ch in positions else ""
            mm = col_meaning(hdr, positions[cm]).lower() if cm in positions else ""
            if "hombres" not in mh or "mujeres" not in mm:
                problems.append(f"{ch}/{cm} sin Hombres/Mujeres")
            age_label = mh.replace("hombres", "").strip(" |")
            age_label = age_label.split(" | ")[-1] if age_label else ""
            ages.append(age_label)
            if sheet != "P2":
                expected = rf"(^|\D)0*{5 * k}\s*(-|a)\s*{5 * k + 4}(\D|$)" if k < 16 else r"80\s*y"
                if not re.search(expected, mh):
                    problems.append(f"{ch} no corresponde al grupo {AGE_GROUPS[k]}")
        rec["age_columns"] = "; ".join(ages)
    rec["layout_check"] = "ok" if not problems else "; ".join(sorted(set(problems)))
    return rec


def run_dictionary_check(cat: pd.DataFrame, years: list[int]) -> pd.DataFrame:
    recs = []
    for r in cat.itertuples():
        for y in years:
            recs.append(check_code(r.series, y, r.module, r.code, r.year_start <= y <= r.year_end))
    df = pd.DataFrame(recs).merge(cat[["module", "code", "indicator", "era"]].rename(columns={"module": "sheet"}), on=["sheet", "code"], how="left")
    return df[["year", "series", "sheet", "code", "in_era", "era", "indicator", "found_in_dictionary", "label", "section", "col01_meaning", "col02_meaning",
               "col03_meaning", "col04_meaning", "col37_meaning", "n_col_tokens", "age_columns", "layout_check", "dictionary_file", "excel_row"]]


# ---------------------------------------------------------------------------
# 3. Lectura de los archivos REM
# ---------------------------------------------------------------------------
def count_lines(path: Path) -> int:
    n = 0
    with open(path, "rb") as fh:
        while chunk := fh.read(1 << 24):
            n += chunk.count(b"\n")
    return n


def read_series(series: str, years: list[int], codes: set[str], chunksize: int, run_log: list[dict]) -> pd.DataFrame:
    root = CFG.PATHS["rem_a"] if series == "A" else CFG.PATHS["rem_p"]
    wanted = set(codes) | {c.lstrip("0") for c in codes if c[0].isdigit()}
    parts = []
    for y in years:
        path = root / f"Serie{series}_{y}.csv"
        if not path.is_file():
            raise FileNotFoundError(f"Falta {path}")
        t0 = time.time()
        n_sel = 0
        for chunk in C.iter_rem_series(path, wanted, chunksize=chunksize):
            chunk = chunk.copy()
            chunk["CodigoPrestacion"] = chunk.CodigoPrestacion.map(lambda v: _norm_code(v, series))
            chunk = chunk.loc[chunk.CodigoPrestacion.isin(codes)]
            chunk["source_file"] = path.name
            n_sel += len(chunk)
            parts.append(chunk)
        seconds = time.time() - t0
        n_lines = count_lines(path) - 1
        st = path.stat()
        run_log.append(dict(series=series, year=y, file=str(path), bytes=st.st_size, mtime_utc=dt.datetime.fromtimestamp(st.st_mtime, dt.UTC).isoformat(),
                            rows_in_file=n_lines, rows_selected=n_sel, seconds=round(seconds, 2)))
        log(f"Serie {series} {y}: {n_lines:,} filas leídas, {n_sel:,} seleccionadas en {seconds:.1f} s")
    df = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(columns=KEYS + RAW_COLS + ["source_file"])
    df["series"] = series
    return df


# ---------------------------------------------------------------------------
# 4. Tabla tidy
# ---------------------------------------------------------------------------
def build_tidy(raw: pd.DataFrame, cat: pd.DataFrame) -> pd.DataFrame:
    df = raw.copy()
    df["year"] = pd.to_numeric(df.Ano, errors="coerce").astype("Int64")
    df["month"] = pd.to_numeric(df.Mes, errors="coerce").astype("Int64")
    if df.year.isna().any() or df.month.isna().any():
        raise ValueError("Año o mes no numérico en REM")
    df["file_year"] = df.source_file.str.extract(r"(\d{4})", expand=False).astype(int)
    if (df.year.astype(int) != df.file_year).any():
        raise ValueError("Columna Ano distinta del año del archivo")
    df["year"] = df.year.astype(int)
    df["month"] = df.month.astype(int)
    meta = cat.set_index("code")
    df["module"] = df.CodigoPrestacion.map(meta.module)
    # Serie P pierde el cero inicial de IdComuna en algunos años (4 dígitos); Serie A siempre trae 5. Se conserva el
    # texto original y se añade la versión armonizada de cinco dígitos (DEIS) para comparar entre series.
    df["id_comuna_5d"] = df.IdComuna.fillna("").str.strip().map(lambda v: v.zfill(5) if v.isdigit() else v)
    df["in_era"] = [meta.year_start[c] <= y <= meta.year_end[c] for c, y in zip(df.CodigoPrestacion, df.year)]
    # Duplicados exactos (misma clave y mismas 50 columnas) y conflictos de clave (misma clave, valores distintos).
    df["exact_duplicate"] = df.duplicated(KEYS + RAW_COLS, keep="first")
    distinct = df.drop_duplicates(KEYS + RAW_COLS)
    conflict_keys = distinct.loc[distinct.duplicated(KEYS, keep=False), KEYS].drop_duplicates()
    conflict_index = pd.MultiIndex.from_frame(conflict_keys)
    df["key_conflict"] = pd.MultiIndex.from_frame(df[KEYS]).isin(conflict_index)
    # Valores numéricos: celda vacía = NaN (distinto de cero); no numéricos se cuentan y quedan como NaN.
    raw_text = df[RAW_COLS]
    num = raw_text.apply(lambda s: pd.to_numeric(s.str.strip().str.replace(",", ".", regex=False), errors="coerce"))
    nonempty = raw_text.apply(lambda s: s.fillna("").str.strip() != "")
    df["n_nonnumeric_cells"] = (nonempty & num.isna()).sum(axis=1)
    df["n_negative_cells"] = (num < 0).sum(axis=1)
    df["col01_num"] = num.Col01
    df["col02_num"] = num.Col02
    df["col03_num"] = num.Col03
    is_a03 = df.module.eq("A03")
    df["total_rule"] = np.where(is_a03, "COL01+COL02 (celdas presentes)", "COL01")
    df["total_known"] = num.Col01
    df.loc[is_a03, "total_known"] = num.loc[is_a03, ["Col01", "Col02"]].sum(axis=1, min_count=1)
    df["a03_one_sex_cell_empty"] = is_a03 & (num.Col01.isna() != num.Col02.isna())
    df["state"] = np.select([df.total_known.isna(), df.total_known < 0, df.total_known == 0], ["reported_empty", "reported_negative", "reported_zero"], "reported_value")
    for raw_c, tidy_c in zip(RAW_COLS, TIDY_COLS):
        df[tidy_c] = df[raw_c].fillna("")
    out = df[["year", "series", "module", "CodigoPrestacion", "IdServicio", "IdRegion", "IdComuna", "id_comuna_5d", "IdEstablecimiento", "month", "in_era"] + TIDY_COLS +
             ["col01_num", "col02_num", "col03_num", "total_rule", "total_known", "state", "a03_one_sex_cell_empty", "exact_duplicate", "key_conflict",
              "n_nonnumeric_cells", "n_negative_cells", "source_file"]].rename(columns={"CodigoPrestacion": "code"})
    out = out.sort_values(["series", "year", "module", "code", "IdEstablecimiento", "month"], kind="mergesort").reset_index(drop=True)
    return out


# ---------------------------------------------------------------------------
# 5. Agregados anuales
# ---------------------------------------------------------------------------
def family_definitions() -> list[dict]:
    """Filas por variante para A05 (ingresos/egresos) y P6 (APS/especialidad); P2 y el resto son idénticos entre variantes."""
    fams = []
    for flow, entry_map, broad in (("entry", CFG.A05_ENTRY, CFG.A05_BROAD_PRE2021["entry"]), ("exit", CFG.A05_EXIT, CFG.A05_BROAD_PRE2021["exit"])):
        base = dict(module="A05", series="A", domain=flow, setting="mental_health", unit=("entries reported" if flow == "entry" else "exits reported"),
                    flow=flow, measures=["annual_sum"])
        noun = "ingresos" if flow == "entry" else "egresos (altas clínicas)"
        for vname, v in CFG.VARIANTS.items():
            fams.append(dict(base, variant=vname, codes=list(v[f"a05_{flow}"]), year_start=2021, year_end=2025,
                             indicator=f"A05 {noun} a salud mental: familia TGD ({v['label']['es']})", identical_across_variants=False))
        fams.append(dict(base, variant="strict_autism", codes=[entry_map["autism"]], year_start=2021, year_end=2025,
                         indicator=f"A05 {noun} a salud mental: autismo estricto", identical_across_variants=True))
        fams.append(dict(base, variant="broad_pre2021", codes=[broad], year_start=2019, year_end=2020,
                         indicator=f"A05 {noun} a salud mental: trastornos generalizados del desarrollo (TGD amplio, pre-2021)", identical_across_variants=True))
    for setting, code_map, broad in (("primary_care", CFG.P6_PRIMARY, CFG.P6_BROAD_PRE2021["primary"]), ("specialty", CFG.P6_SPECIALTY, CFG.P6_BROAD_PRE2021["specialty"])):
        key = "p6_primary" if setting == "primary_care" else "p6_specialty"
        base = dict(module="P6", series="P", domain="follow_up", setting=setting, unit="people in stock", flow="stock", measures=["december_stock", "june_stock"])
        lugar = "APS" if setting == "primary_care" else "especialidad"
        for vname, v in CFG.VARIANTS.items():
            fams.append(dict(base, variant=vname, codes=list(v[key]), year_start=2021, year_end=2025,
                             indicator=f"P6 población bajo control en {lugar}: familia TGD ({v['label']['es']})", identical_across_variants=False))
        fams.append(dict(base, variant="strict_autism", codes=[code_map["autism"]], year_start=2021, year_end=2025,
                         indicator=f"P6 población bajo control en {lugar}: autismo estricto", identical_across_variants=True))
        fams.append(dict(base, variant="broad_pre2021", codes=[broad], year_start=2019, year_end=2020,
                         indicator=f"P6 población bajo control en {lugar}: TGD amplio (pre-2021)", identical_across_variants=True))
    return fams


MEASURE_MONTHS = {"annual_sum": None, "december_stock": 12, "june_stock": 6}


def summarise_group(rows: pd.DataFrame, years: list[int], measure: str) -> tuple[list[dict], set]:
    """Resumen anual de un conjunto de filas tidy (un código o una familia); devuelve filas y el panel estable."""
    month = MEASURE_MONTHS[measure]
    sub = rows if month is None else rows.loc[rows.month == month]
    stable = set.intersection(*[set(sub.loc[sub.year == y, "IdEstablecimiento"]) for y in years]) if years else set()
    out = []
    for y in years:
        r = sub.loc[sub.year == y]
        nodup = r.loc[~r.exact_duplicate]
        in_panel = r.loc[r.IdEstablecimiento.isin(stable)]
        out.append(dict(
            year=y, measure=measure, month=(f"{month:02d}" if month else "01-12"),
            total=r.total_known.sum(min_count=1), total_dedup=nodup.total_known.sum(min_count=1),
            n_reporting_establishments=int(r.IdEstablecimiento.nunique()),
            n_establishments_value_gt0=int(r.loc[r.total_known > 0, "IdEstablecimiento"].nunique()),
            n_rows=int(len(r)), n_rows_value=int((r.state == "reported_value").sum()), n_rows_zero=int((r.state == "reported_zero").sum()),
            n_rows_empty=int((r.state == "reported_empty").sum()), n_rows_exact_duplicate=int(r.exact_duplicate.sum()),
            months_covered=int(r.month.nunique()), n_stable_panel_establishments=len(stable),
            stable_panel_total=in_panel.total_known.sum(min_count=1),
            note="" if len(r) else "no rows reported for this code-year",
        ))
    return out, stable


def build_annual(tidy: pd.DataFrame, cat: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    era_rows = tidy.loc[tidy.in_era]
    out, panels = [], {}
    for r in cat.itertuples():
        rows = era_rows.loc[era_rows.code == r.code]
        years = list(range(r.year_start, r.year_end + 1))
        measures = ["annual_sum"] if r.series == "A" else ["december_stock", "june_stock"]
        for m in measures:
            recs, stable = summarise_group(rows, years, m)
            panels[(r.code, m)] = stable
            for rec in recs:
                out.append(dict(series=r.series, module=r.module, code=r.code, variant="single_code", identical_across_variants=True, indicator=r.indicator,
                                era=r.era, era_start=r.year_start, era_end=r.year_end, domain=r.domain, setting=r.setting, unit=r.unit,
                                aggregation_rule=("sum of months; " + ("COL01+COL02 (cells present)" if r.module == "A03" else "COL01")) if m == "annual_sum"
                                else f"stock MES={MEASURE_MONTHS[m]:02d}; COL01; never summed with the other semester",
                                comparability_warning=r.comparability_warning, **rec))
    for f in family_definitions():
        rows = era_rows.loc[era_rows.code.isin(f["codes"])]
        years = list(range(f["year_start"], f["year_end"] + 1))
        for m in f["measures"]:
            recs, stable = summarise_group(rows, years, m)
            panels[("+".join(f["codes"]), m)] = stable
            for rec in recs:
                out.append(dict(series=f["series"], module=f["module"], code="+".join(f["codes"]), variant=f["variant"], identical_across_variants=f["identical_across_variants"],
                                indicator=f["indicator"], era=f"{f['year_start']}–{f['year_end']}", era_start=f["year_start"], era_end=f["year_end"], domain=f["domain"],
                                setting=f["setting"], unit=f["unit"],
                                aggregation_rule=("sum of months and of the family codes; COL01" if m == "annual_sum" else f"stock MES={MEASURE_MONTHS[m]:02d}; COL01 summed over the family codes; never summed with the other semester"),
                                comparability_warning=("Family sum over separately reported categories; establishments counted once if they report any family code"
                                                       + ("; broad pre-2021 category cannot isolate autism" if f["variant"] == "broad_pre2021" else "")),
                                **rec))
    ann = pd.DataFrame(out)
    ann["script"] = SCRIPT
    ann["source_file"] = np.where(ann.series == "A", "REM/SerieA/SerieA_{year}.csv", "REM/SerieP/SerieP_{year}.csv")
    ann["source_file"] = [s.format(year=y) for s, y in zip(ann.source_file, ann.year)]
    cols = ["year", "series", "module", "code", "variant", "identical_across_variants", "indicator", "era", "era_start", "era_end", "domain", "setting", "unit",
            "measure", "month", "aggregation_rule", "total", "total_dedup", "n_reporting_establishments", "n_establishments_value_gt0", "n_rows", "n_rows_value",
            "n_rows_zero", "n_rows_empty", "n_rows_exact_duplicate", "months_covered", "n_stable_panel_establishments", "stable_panel_total",
            "comparability_warning", "note", "source_file", "script"]
    ann = ann[cols].sort_values(["series", "module", "variant", "code", "measure", "year"], kind="mergesort").reset_index(drop=True)
    return ann, panels


# ---------------------------------------------------------------------------
# 6. A05 por edad y sexo (COL02/COL03 = sexo; COL04..COL37 = 17 grupos × sexo)
# ---------------------------------------------------------------------------
def build_a05_age_sex(tidy: pd.DataFrame, cat: pd.DataFrame) -> pd.DataFrame:
    a05 = tidy.loc[tidy.in_era & tidy.module.eq("A05")].copy()
    cells = {}
    for k, age in enumerate(AGE_GROUPS):
        cells[("H", age)] = f"col{4 + 2 * k:02d}"
        cells[("M", age)] = f"col{5 + 2 * k:02d}"
    cells[("H", "total")] = "col02"
    cells[("M", "total")] = "col03"
    num = a05[list(dict.fromkeys(cells.values()))].apply(lambda s: pd.to_numeric(s.replace("", np.nan), errors="coerce"))
    num["year"] = a05.year.values
    num["code"] = a05.code.values
    num["est"] = a05.IdEstablecimiento.values  # establecimientos reportantes (≥1 fila del código en el año), exigidos en toda tabla REM
    category = {v: k for k, v in CFG.A05_ENTRY.items()} | {v: k for k, v in CFG.A05_EXIT.items()} | {CFG.A05_BROAD_PRE2021["entry"]: "broad_pdd", CFG.A05_BROAD_PRE2021["exit"]: "broad_pdd"}
    flow = {v: "entry" for v in CFG.A05_ENTRY.values()} | {v: "exit" for v in CFG.A05_EXIT.values()} | {CFG.A05_BROAD_PRE2021["entry"]: "entry", CFG.A05_BROAD_PRE2021["exit"]: "exit"}
    label = cat.set_index("code").indicator
    recs = []
    for (y, code), g in num.groupby(["year", "code"]):
        n_est = int(g.est.nunique())
        for (sex, age), col in cells.items():
            recs.append(dict(year=y, module="A05", code=code, variant="single_code", flow=flow[code], category=category[code], indicator=label[code],
                             sex=SEX_LABELS[sex], age_group=age, source_col=col.upper(), count=g[col].sum(min_count=1), n_rows=len(g), n_rows_cell_present=int(g[col].notna().sum()),
                             n_reporting_establishments=n_est))
    single = pd.DataFrame(recs)
    fams = []
    for fname, v in CFG.VARIANTS.items():
        for fl in ("entry", "exit"):
            codes = v[f"a05_{fl}"]
            s = single.loc[single.code.isin(codes)]
            g = s.groupby(["year", "sex", "age_group", "source_col"], as_index=False).agg(count=("count", lambda x: x.sum(min_count=1)), n_rows=("n_rows", "sum"), n_rows_cell_present=("n_rows_cell_present", "sum"))
            # Establecimientos contados una sola vez si reportan cualquier código de la familia en el año (no se suman por código).
            fam_est = num.loc[num.code.isin(codes)].groupby("year").est.nunique()
            g["n_reporting_establishments"] = g.year.map(fam_est).fillna(0).astype(int)
            g["module"], g["code"], g["variant"], g["flow"], g["category"] = "A05", "+".join(codes), fname, fl, "pdd_family"
            g["indicator"] = f"A05 {'ingresos' if fl == 'entry' else 'egresos'}: familia TGD ({v['label']['es']})"
            fams.append(g)
    df = pd.concat([single] + fams, ignore_index=True)
    df["unit"] = np.where(df.flow == "entry", "entries reported", "exits reported")
    df["age_group_order"] = df.age_group.map({a: i for i, a in enumerate(AGE_GROUPS + ["total"])})
    df = df.sort_values(["variant", "flow", "code", "year", "sex", "age_group_order"], kind="mergesort").drop(columns="age_group_order").reset_index(drop=True)
    return df[["year", "module", "code", "variant", "flow", "category", "indicator", "sex", "age_group", "source_col", "unit", "count", "n_rows", "n_rows_cell_present",
               "n_reporting_establishments"]]


# ---------------------------------------------------------------------------
# 7. Panel establecimiento-año
# ---------------------------------------------------------------------------
def build_establishment_year(tidy: pd.DataFrame, cat: pd.DataFrame, panels: dict) -> pd.DataFrame:
    era_rows = tidy.loc[tidy.in_era].copy()
    geo_check = era_rows.groupby(["year", "IdEstablecimiento"])[["IdServicio", "IdRegion", "id_comuna_5d"]].nunique()
    n_geo_conflicts = int((geo_check > 1).any(axis=1).sum())
    era_rows["is_value"] = era_rows.state.eq("reported_value")
    era_rows["is_zero"] = era_rows.state.eq("reported_zero")
    era_rows["is_empty"] = era_rows.state.eq("reported_empty")
    era_rows["dec_value"] = era_rows.total_known.where(era_rows.month == 12)
    era_rows["jun_value"] = era_rows.total_known.where(era_rows.month == 6)
    era_rows["has_dec"] = era_rows.month == 12
    era_rows["has_jun"] = era_rows.month == 6
    g = era_rows.groupby(["year", "series", "module", "code", "IdEstablecimiento"], as_index=False).agg(
        IdServicio=("IdServicio", "first"), IdRegion=("IdRegion", "first"), IdComuna=("IdComuna", "first"), id_comuna_5d=("id_comuna_5d", "first"),
        months_with_rows=("month", "nunique"), n_rows=("month", "size"), months_value=("is_value", "sum"), months_zero=("is_zero", "sum"),
        months_empty=("is_empty", "sum"), n_exact_duplicate_rows=("exact_duplicate", "sum"),
        annual_total=("total_known", lambda s: s.sum(min_count=1)), december_value=("dec_value", lambda s: s.sum(min_count=1)),
        june_value=("jun_value", lambda s: s.sum(min_count=1)), has_december_row=("has_dec", "any"), has_june_row=("has_jun", "any"))
    is_p = g.series.eq("P")
    g[["has_december_row", "has_june_row"]] = g[["has_december_row", "has_june_row"]].astype("boolean")
    g.loc[is_p, "annual_total"] = np.nan  # stock: nunca sumar semestres
    g.loc[~is_p, ["december_value", "june_value"]] = np.nan
    g.loc[~is_p, ["has_december_row", "has_june_row"]] = pd.NA
    g["in_stable_panel"] = [est in panels.get((code, "annual_sum" if s == "A" else "december_stock"), set()) for code, s, est in zip(g.code, g.series, g.IdEstablecimiento)]
    g["value_rule"] = np.where(is_p, "december_value = COL01 in MES=12; june_value = COL01 in MES=06", np.where(g.module.eq("A03"), "annual_total = sum of months of COL01+COL02", "annual_total = sum of months of COL01"))
    g["n_geo_conflicts_year_establishment"] = n_geo_conflicts
    g = g.sort_values(["series", "module", "code", "year", "IdEstablecimiento"], kind="mergesort").reset_index(drop=True)
    return g[["year", "series", "module", "code", "IdEstablecimiento", "IdServicio", "IdRegion", "IdComuna", "id_comuna_5d", "months_with_rows", "n_rows", "months_value", "months_zero",
              "months_empty", "n_exact_duplicate_rows", "annual_total", "december_value", "june_value", "has_december_row", "has_june_row", "in_stable_panel", "value_rule",
              "n_geo_conflicts_year_establishment"]]


# ---------------------------------------------------------------------------
# 8. Controles de reproducción
# ---------------------------------------------------------------------------
def build_controls(ann: pd.DataFrame, tidy: pd.DataFrame, agesex: pd.DataFrame, dchk: pd.DataFrame, run_log: list[dict]) -> pd.DataFrame:
    single = ann.loc[ann.variant == "single_code"].set_index(["code", "measure", "year"])

    def value(code, measure, year, field="total"):
        try:
            return single.loc[(code, measure, year), field]
        except KeyError:
            return np.nan

    rows = []

    def add(name, key, expected, observed, note="", kind="count"):
        exp = np.nan if expected is None else float(expected)
        obs = np.nan if observed is None or (isinstance(observed, float) and np.isnan(observed)) else float(observed)
        if np.isnan(exp):
            status, absd, reld = "info", np.nan, np.nan
        elif np.isnan(obs):
            status, absd, reld = "differs", np.nan, np.nan
            note = (note + "; " if note else "") + "observed value missing"
        else:
            absd = obs - exp
            reld = absd / exp if exp else np.nan
            status = "ok" if (absd == 0 if kind == "count" else abs(reld) <= 0.005) else "differs"
        rows.append(dict(name=name, key=key, expected=exp, observed=obs, abs_diff=absd, rel_diff=reld, status=status, note=note))

    CT = CFG.CONTROLS
    simple = {"a05_autism_entries": ("05990022", "annual_sum"), "a05_autism_exits": ("05990027", "annual_sum"),
              "a27_counselling": ("29101566", "annual_sum"), "a27_assisted_referral": ("29101574", "annual_sum"),
              "a28_primary": ("29101629", "annual_sum"), "a28_hospital": ("29101651", "annual_sum"),
              "p2_tea_december": ("P2500500", "december_stock"), "p2_naneas_total_december": ("P2501878", "december_stock"),
              "a03_legacy_mchat_done": ("03500406", "annual_sum"), "a03_legacy_mchat_altered": ("03500407", "annual_sum")}
    for name, (code, measure) in simple.items():
        for y, exp in CT[name].items():
            add(name, f"{code}|{measure}|{y}", exp, value(code, measure, y), note="config.CONTROLS")
    for y, exp in CT["p2_establishments_december"].items():
        add("p2_establishments_december", f"P2500500|december_stock|{y}", exp, value("P2500500", "december_stock", y, "n_reporting_establishments"), note="config.CONTROLS; distinct IdEstablecimiento with a December row")
    for name, pre, post in (("p6_primary_december", "P6223000", "P6241010"), ("p6_specialty_december", "P6223380", "P6241060")):
        for y, exp in CT[name].items():
            code = pre if y <= 2020 else post
            add(name, f"{code}|december_stock|{y}", exp, value(code, "december_stock", y), note="config.CONTROLS; broad TGD 2019–2020, strict autism 2021+ (definition break)")
    for y, (lo, me, hi) in CT["a03_2023_low_medium_high"].items():
        for code, exp, lab in (("09600213", lo, "low"), ("09600214", me, "medium"), ("09600215", hi, "high")):
            add(f"a03_2023_2024_{lab}", f"{code}|annual_sum|{y}", exp, value(code, "annual_sum", y), note="config.CONTROLS; COL01+COL02 cells present")
    for name, spec in EXTRA_CONTROLS.items():
        for y, exp in spec["values"].items():
            add(name, f"{spec['code']}|{spec['measure']}|{y}", exp, value(spec["code"], spec["measure"], y, spec["field"]), note="scripts/downloads/DATA_REVIEW.md")
    # Informativos solicitados: establecimientos A05 autismo por año y paneles estables.
    for y in range(2021, 2026):
        add("a05_autism_entries_establishments", f"05990022|annual_sum|{y}", None, value("05990022", "annual_sum", y, "n_reporting_establishments"), note="distinct IdEstablecimiento with ≥1 row in the year")
    key_series = [("05990022", "annual_sum"), ("05990027", "annual_sum"), ("P2500500", "december_stock"), ("P6241010", "december_stock"), ("P6241060", "december_stock"),
                  ("29101566", "annual_sum"), ("29101574", "annual_sum"), ("29101629", "annual_sum"), ("29101651", "annual_sum"), ("03500406", "annual_sum"), ("03500407", "annual_sum")]
    for code, measure in key_series:
        s = single.xs((code, measure), level=("code", "measure"))
        for y, r in s.iterrows():
            add("stable_panel_total", f"{code}|{measure}|{y}", None, r.stable_panel_total, note=f"{int(r.n_stable_panel_establishments)} establishments reporting in every year of the era; observed-panel total {r.total:.0f}")
    fam = ann.loc[ann.variant.isin(["con_rett", "sin_rett"]) & ann.measure.isin(["annual_sum", "december_stock"])]
    for r in fam.itertuples():
        add(f"variant_total_{r.variant}", f"{r.module}|{r.setting}|{r.code}|{r.measure}|{r.year}", None, r.total, note=f"{r.indicator}; establishments {r.n_reporting_establishments}")
    # Calidad de datos.
    add("data_quality_exact_duplicate_rows", "SerieA", None, int(tidy.loc[tidy.series == "A", "exact_duplicate"].sum()), note="kept in totals (reproduces controls); see total_dedup")
    add("data_quality_exact_duplicate_rows", "SerieP", None, int(tidy.loc[tidy.series == "P", "exact_duplicate"].sum()), note="kept in totals; see total_dedup")
    add("data_quality_key_conflict_rows", "all", None, int(tidy.key_conflict.sum()), note="same establishment×month×code with different values")
    add("data_quality_nonnumeric_cells", "all", None, int(tidy.n_nonnumeric_cells.sum()))
    add("data_quality_negative_cells", "all", None, int(tidy.n_negative_cells.sum()))
    add("data_quality_out_of_era_rows", "all", None, int((~tidy.in_era).sum()), note="rows of a code outside its rem_pathway_codes.csv era (excluded from annual tables)")
    add("data_quality_serieP_months_not_6_12", "all", None, int((~tidy.loc[tidy.series == "P", "month"].isin([6, 12])).sum()))
    p4 = tidy.loc[(tidy.series == "P") & (tidy.IdComuna.str.len() == 4)]
    add("data_quality_serieP_idcomuna_4_digits", "SerieP", None, int(len(p4)), note=f"rows whose IdComuna lost the leading zero (years {sorted(p4.year.unique().tolist())}); harmonised in id_comuna_5d")
    geo = tidy.loc[tidy.in_era].groupby(["year", "IdEstablecimiento"])[["IdServicio", "IdRegion", "id_comuna_5d"]].nunique()
    add("data_quality_establishment_geography_conflicts", "all", None, int((geo > 1).any(axis=1).sum()), note="establishment-years with >1 IdServicio/IdRegion/id_comuna_5d across both series; expected 0")
    add("data_quality_a03_rows_one_sex_cell_empty", "A03", None, int(tidy.a03_one_sex_cell_empty.sum()), note=f"of {int(tidy.module.eq('A03').sum())} A03 rows; summed with cells present")
    a05 = tidy.loc[tidy.in_era & tidy.module.eq("A05")]
    agecols = [f"col{i:02d}" for i in range(4, 38)]
    agenum = a05[agecols].apply(lambda s: pd.to_numeric(s.replace("", np.nan), errors="coerce"))
    add("data_quality_a05_sex_sum_mismatch", "A05", None, int(((a05.col02_num.fillna(0) + a05.col03_num.fillna(0)) != a05.col01_num.fillna(0)).sum()), note="rows where COL02+COL03 != COL01")
    add("data_quality_a05_age_sum_mismatch", "A05", None, int((agenum.fillna(0).sum(axis=1) != a05.col01_num.fillna(0)).sum()), note="rows where sum(COL04..COL37) != COL01")
    add("dictionary_codes_found_in_era", "all", None, int((dchk.in_era & dchk.found_in_dictionary).sum()), note=f"of {int(dchk.in_era.sum())} code-years in era")
    add("dictionary_codes_found_outside_era", "all", None, int((~dchk.in_era & dchk.found_in_dictionary).sum()), note="expected 0")
    add("dictionary_layout_problems", "all", None, int((dchk.in_era & dchk.found_in_dictionary & (dchk.layout_check != "ok")).sum()), note="expected 0")
    for r in run_log:
        add("runtime_seconds", f"Serie{r['series']}_{r['year']}.csv", None, r["seconds"], note=f"rows_in_file={r['rows_in_file']}; rows_selected={r['rows_selected']}")
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--years", type=int, nargs=2, metavar=("FIRST", "LAST"), default=(CFG.YEARS_REM[0], CFG.YEARS_REM[-1]))
    ap.add_argument("--chunksize", type=int, default=1_000_000)
    args = ap.parse_args(argv)
    years = list(range(args.years[0], args.years[1] + 1))
    t_start = time.time()
    cat = load_catalogue()
    log(f"{len(cat)} códigos en el catálogo ({(cat.series == 'A').sum()} Serie A, {(cat.series == 'P').sum()} Serie P)")

    t0 = time.time()
    dchk = run_dictionary_check(cat, years)
    log(f"diccionarios verificados: {int((dchk.in_era & dchk.found_in_dictionary).sum())}/{int(dchk.in_era.sum())} código-año en era encontrados en {time.time() - t0:.1f} s")

    run_log: list[dict] = []
    raw_a = read_series("A", years, set(cat.loc[cat.series == "A", "code"]), args.chunksize, run_log)
    raw_p = read_series("P", years, set(cat.loc[cat.series == "P", "code"]), args.chunksize, run_log)
    tidy = build_tidy(pd.concat([raw_a, raw_p], ignore_index=True), cat)
    log(f"tidy: {len(tidy):,} filas; estados {tidy.state.value_counts().to_dict()}")

    ann, panels = build_annual(tidy, cat)
    agesex = build_a05_age_sex(tidy, cat)
    est = build_establishment_year(tidy, cat, panels)
    controls = build_controls(ann, tidy, agesex, dchk, run_log)

    C.atomic_write_csv(dchk, CFG.TIDY / "rem_code_dictionary_check.csv")
    C.atomic_write_csv(tidy, CFG.TIDY / "rem_pathway_tidy.csv")
    C.atomic_write_csv(ann, CFG.TIDY / "rem_pathway_annual.csv")
    C.atomic_write_csv(agesex, CFG.TIDY / "rem_a05_age_sex_annual.csv")
    C.atomic_write_csv(est, CFG.TIDY / "rem_establishment_year.csv")
    C.atomic_write_csv(controls, CONTROLS_DIR / f"{MODULE}_controls.csv")
    total_seconds = round(time.time() - t_start, 1)
    C.atomic_write_json(dict(module=MODULE, script=SCRIPT, run_utc=dt.datetime.now(dt.UTC).isoformat(timespec="seconds"), years=years, total_seconds=total_seconds,
                             files=run_log, rows_tidy=len(tidy), rows_annual=len(ann), controls_ok=int((controls.status == "ok").sum()),
                             controls_differs=int((controls.status == "differs").sum())), CONTROLS_DIR / f"{MODULE}_run_log.json")
    n_ok, n_diff = int((controls.status == "ok").sum()), int((controls.status == "differs").sum())
    log(f"controles: {n_ok} ok, {n_diff} difieren; tiempo total {total_seconds} s")
    if n_diff:
        print(controls.loc[controls.status == "differs"].to_string(), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())

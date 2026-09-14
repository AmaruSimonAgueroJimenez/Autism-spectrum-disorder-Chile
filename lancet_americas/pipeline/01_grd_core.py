#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""01_grd_core.py — Núcleo hospitalario GRD 2019–2024 (paso 01 del pipeline Lancet Regional Health – Americas).

Lee GRD_PUBLICO_{2019..2024}.csv (sep='|', todo como texto) por trozos de 200.000 filas mediante
`common.iter_grd`, conserva en memoria solo las filas con F84 y agregados por trozo, y escribe (escritura
atómica, nunca microdatos):

  outputs/tidy/grd_year_summary.csv             año × variante × panel × actividad × posición: episodios con F84,
                                                denominador de episodios del mismo panel/actividad, tasa por 100.000
                                                episodios con IC exacto de Poisson, personas únicas dentro del año,
                                                profundidad diagnóstica (media/mediana) de todos los episodios y de los F84.
  outputs/tidy/grd_hospital_year.csv            año × variante × hospital (nombre desde el catálogo GRD).
  outputs/tidy/grd_age_sex_year.csv             año × variante × panel × posición × actividad × sexo × grupo de edad,
                                                numerador y denominador compatibles (misma fuente, panel y actividad).
  outputs/tidy/grd_coding_depth_year.csv        año × variante × panel × actividad × tramo de profundidad diagnóstica.
  outputs/tidy/grd_subcode_year.csv             año × subcódigo F84 × posición (un registro cuenta una vez por subcódigo y posición).
  outputs/tidy/grd_identifier_audit.csv         año × variante: identificadores válidos, únicos, longitud modal y solapamiento
                                                con el año anterior (informativo; nunca se deduplica entre años).
  outputs/tidy/grd_activity_categories_year.csv año × categoría cruda de TIPO_ACTIVIDAD × clase analítica.
  outputs/tidy/grd_fixed_panel_hospitals.csv    panel fijo (hospitales presentes en todos los años 2019–2022) y presencia anual.
  outputs/controls/01_grd_core_controls.csv     esperado (config.CONTROLS) frente a observado.
  outputs/controls/01_grd_core_runlog.json      procedencia (SHA-256, bytes, mtime), tiempos por año, valores crudos observados.

Definiciones operativas (diccionario de variables en el informe estructurado):
  * episodio = una fila del GRD publicado; no se eliminan filas. Los duplicados exactos sobre las columnas analizadas
    se cuentan e informan (antes/después) sin aplicarse, porque los controles se reproducen sin eliminación.
  * profundidad diagnóstica = número de campos DIAGNOSTICO1..35 no vacíos tras `common.normalize_code`
    (sin puntos ni espacios, mayúsculas); cadena vacía o NA = ausente.
  * F84 = pertenencia exacta del código normalizado a `config.VARIANTS[v]['grd_subcodes']`:
    con_rett = {F84, F840, F841, F842, F843, F844, F845, F848, F849}; sin_rett = con_rett − {F842};
    strict_autism_f840 = {F840} (serie extra idéntica entre variantes).
  * posición: principal = DIAGNOSTICO1; secundaria = DIAGNOSTICO2..35; any = principal o secundaria;
    secondary_only = secundaria y no principal; principal_and_secondary = ambas.
  * actividad: hospitalisation = TIPO_ACTIVIDAD == 'HOSPITALIZACIÓN' (comparación sin acentos, mayúsculas, sin
    espacios extremos); cma = 'CIRUGÍA MAYOR AMBULATORIA (CMA)'; other = resto (solo existe en 2019:
    hospitalización en urgencia, hospitalización diurna, desconocido, no identificado); all = sin filtro.
  * panel: observed = hospitales con ≥1 episodio en el año; fixed65 = intersección de los hospitales presentes en
    2019, 2020, 2021 y 2022 (se verifica que sean 65 y subconjunto de 2023 y 2024).
  * personas únicas = identificadores válidos distintos dentro del año y de la celda (CIP_ENCRIPTADO 2019–2023,
    ID_BENEFICIARIO 2024). Nunca se deduplica a través de 2020/2021 ni entre años.
  * edad = floor((FECHA_INGRESO − FECHA_NACIMIENTO) / 365,25) años; <0 o >110 → desconocida; grupos quinquenales
    OMS 0-4 … 75-79, 80+ (`epi_helpers.AGE_GROUPS`). Fechas ISO (AAAA-MM-DD) o DD-MM-AAAA (2023), nunca inferidas.

Uso: python3 lancet_americas/pipeline/01_grd_core.py [--years 2019 ... 2024] [--chunksize 200000] [--no-hash]
     [--max-chunks N (solo pruebas: los controles no se reproducen)]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

LANCET_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(LANCET_DIR))
import config as CFG  # noqa: E402
import common as C  # noqa: E402

MODULE = "01_grd_core"
SCRIPT = str(Path(__file__).resolve().relative_to(CFG.REPO))
# config.CONTROLS es el diccionario de valores esperados (sombrea la ruta); la carpeta se deriva de config.OUT.
CONTROLS_DIR = CFG.OUT / "controls"
CONTROLS_DIR.mkdir(parents=True, exist_ok=True)

DIAGS = C.GRD_DIAG_COLS
SECONDARY_DIAGS = DIAGS[1:]
ID_CANDIDATES = ["CIP_ENCRIPTADO", "ID_BENEFICIARIO"]
BASE_COLS = ["COD_HOSPITAL", "SEXO", "FECHA_NACIMIENTO", "FECHA_INGRESO", "FECHAALTA", "COMUNA", "REGION",
             "TIPO_ACTIVIDAD", "TIPO_INGRESO", "TIPOALTA"]

VARIANT_CODES: dict[str, set[str]] = {name: set(spec["grd_subcodes"]) for name, spec in CFG.VARIANTS.items()}
VARIANT_CODES["strict_autism_f840"] = {"F840"}
VARIANTS = list(VARIANT_CODES)
PRIMARY_VARIANT = "con_rett"  # los controles del brief corresponden a la familia F84 completa

ACTIVITIES = ["all", "hospitalisation", "cma", "other"]
POSITIONS = ["any", "principal", "secondary_only", "principal_and_secondary"]
PANELS = ["observed", "fixed65"]
FIXED_PANEL_YEARS = [2019, 2020, 2021, 2022]
SEX_LEVELS = ["HOMBRE", "MUJER", "unknown"]
AGE_LEVELS = list(C.AGE_GROUPS) + ["unknown"]
DEPTH_BINS = [("0", 0, 0), ("1", 1, 1), ("2", 2, 2), ("3", 3, 3), ("4", 4, 4), ("5", 5, 5),
              ("6-7", 6, 7), ("8-10", 8, 10), ("11+", 11, 10 ** 6)]
# Mismo criterio que scripts/audit_grd_linkage.valid_id (re-implementado para que el módulo sea autocontenido).
INVALID_ID_VALUES = {"", "0", "-1", "NA", "NAN", "NULL", "DESCONOCIDO", "SIN INFORMACION", "SIN INFORMACIÓN"}


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def strip_accents(text: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch))


def normalise_codes(series: pd.Series) -> np.ndarray:
    """Equivalente vectorizado de `common.normalize_code` (se verifica en el primer trozo de cada año)."""
    out = (series.fillna("").astype(object).str.replace(".", "", regex=False).str.replace(" ", "", regex=False)
           .str.strip().str.upper())
    return out.to_numpy(dtype=object)


def classify_activity(raw: str) -> str:
    key = strip_accents(raw or "").upper().strip()
    if key == "HOSPITALIZACION":
        return "hospitalisation"
    if key.startswith("CIRUGIA MAYOR AMBULATORIA"):
        return "cma"
    return "other"


def parse_dates(series: pd.Series) -> pd.Series:
    """Analiza solo los dos formatos observados (AAAA-MM-DD y DD-MM-AAAA); todo lo demás es NaT."""
    s = series.fillna("").astype(object).str.strip()
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    iso = s.str.match(r"^\d{4}-\d{2}-\d{2}$").to_numpy(dtype=bool)
    dmy = s.str.match(r"^\d{2}-\d{2}-\d{4}$").to_numpy(dtype=bool)
    if iso.any():
        out.loc[iso] = pd.to_datetime(s.loc[iso], format="%Y-%m-%d", errors="coerce")
    if dmy.any():
        out.loc[dmy] = pd.to_datetime(s.loc[dmy], format="%d-%m-%Y", errors="coerce")
    return out


def valid_id_mask(ids: pd.Series) -> np.ndarray:
    return (~ids.str.upper().isin(INVALID_ID_VALUES)).to_numpy(dtype=bool)


def weighted_median(values: np.ndarray, counts: np.ndarray) -> float:
    """Mediana de una distribución discreta dada por valores y frecuencias (promedio de los dos centrales si n es par)."""
    order = np.argsort(values)
    v, c = np.asarray(values)[order], np.asarray(counts)[order]
    n = int(c.sum())
    if n == 0:
        return float("nan")
    cum = np.cumsum(c)
    if n % 2 == 1:
        return float(v[np.searchsorted(cum, (n + 1) // 2)])
    lo = v[np.searchsorted(cum, n // 2)]
    hi = v[np.searchsorted(cum, n // 2 + 1)]
    return float((lo + hi) / 2)


def depth_bin(depth: int) -> str:
    for label, lo, hi in DEPTH_BINS:
        if lo <= depth <= hi:
            return label
    return "11+"


def hospital_catalogue() -> tuple[dict[str, str], str]:
    """Catálogo código → nombre desde las tablas maestras GRD (hoja 'Hospitales'); se registra el archivo usado."""
    candidates = [CFG.PATHS["grd_metadata"] / "official" / "tablas_maestras_bases_GRD.xlsx",
                  CFG.PATHS["grd_metadata"] / "códigos grd.xlsx"]
    for path in candidates:
        if not path.is_file():
            continue
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            if "Hospitales" not in wb.sheetnames:
                wb.close()
                continue
            names: dict[str, str] = {}
            for row in wb["Hospitales"].iter_rows(max_col=2, values_only=True):
                if row[0] is not None and re.fullmatch(r"\d{6}", str(row[0]).strip()) and row[1]:
                    names.setdefault(str(row[0]).strip(), str(row[1]).strip())
            wb.close()
            return names, str(path)
        except Exception as exc:  # noqa: BLE001 — el catálogo es opcional; el código de hospital siempre se conserva
            print(f"  aviso: no se pudo leer {path}: {exc}", flush=True)
    return {}, ""


def deis_catalogue() -> tuple[dict[str, str], str]:
    """Respaldo: catastro DEIS vigente (EstablecimientoCodigo → EstablecimientoGlosa) para códigos ausentes del catálogo GRD."""
    path = CFG.PATHS["establecimientos"] / "establecimientos_salud_vigentes.csv"
    if not path.is_file():
        return {}, ""
    try:
        df = pd.read_csv(path, sep=";", dtype=str, usecols=["EstablecimientoCodigo", "EstablecimientoGlosa"],
                         encoding="utf-8", encoding_errors="replace")
        df = df.dropna()
        return dict(zip(df["EstablecimientoCodigo"].str.strip(), df["EstablecimientoGlosa"].str.strip())), str(path)
    except Exception as exc:  # noqa: BLE001
        print(f"  aviso: no se pudo leer {path}: {exc}", flush=True)
        return {}, ""


class HospitalNames:
    """Nombre y fuente del nombre por código: 'grd_master' (tablas maestras GRD), 'deis_catalogue' (catastro DEIS) o 'code'."""

    def __init__(self) -> None:
        self.grd, self.grd_path = hospital_catalogue()
        self.deis, self.deis_path = deis_catalogue()

    def name(self, code: str) -> str:
        return self.grd.get(code) or self.deis.get(code) or code

    def source(self, code: str) -> str:
        return "grd_master" if code in self.grd else ("deis_catalogue" if code in self.deis else "code")


# ---------------------------------------------------------------------------
# Lectura anual
# ---------------------------------------------------------------------------
def verify_full_row_duplicates(path: Path, dup_rows: pd.DataFrame, id_col: str, chunksize: int) -> tuple[int, list[str]]:
    """Segunda pasada dirigida: recupera con las 129 columnas las filas F84 duplicadas sobre las columnas leídas y
    comprueba si son duplicados exactos de fila completa; devuelve el número de duplicados completos y las columnas
    (fuera de la llave) en que difieren las filas de cada grupo."""
    key_cols = ["COD_HOSPITAL", id_col, "FECHA_INGRESO", "FECHAALTA", "DIAGNOSTICO1"]
    keys = dup_rows[key_cols].drop_duplicates()
    parts = []
    for chunk in pd.read_csv(path, sep="|", dtype="string", chunksize=chunksize, encoding="utf-8", encoding_errors="replace"):
        hit = chunk.merge(keys, how="inner", on=key_cols)
        if len(hit):
            parts.append(hit)
    if not parts:
        return 0, []
    full = pd.concat(parts, ignore_index=True)
    n_full = int(full.duplicated(keep="first").sum())
    differing: set[str] = set()
    for _, g in full.groupby(key_cols, dropna=False):
        if len(g) < 2:
            continue
        for col in full.columns:
            if col not in key_cols and g[col].fillna("<NA>").nunique() > 1:
                differing.add(col)
    return n_full, sorted(differing)


def read_year(year: int, chunksize: int, max_chunks: int | None, do_hash: bool) -> dict:
    path = CFG.PATHS["grd"] / f"GRD_PUBLICO_{year}.csv"
    if not path.is_file():
        raise FileNotFoundError(path)
    header = pd.read_csv(path, sep="|", nrows=0).columns.tolist()
    id_cols = [c for c in ID_CANDIDATES if c in header]
    if len(id_cols) != 1:
        raise ValueError(f"{path.name}: se esperaba exactamente una columna identificadora, hay {id_cols}")
    id_col = id_cols[0]
    usecols = ["COD_HOSPITAL", id_col] + [c for c in BASE_COLS if c != "COD_HOSPITAL"] + DIAGS
    present = [c for c in usecols if c in header]
    missing = sorted(set(usecols) - set(present))

    t0 = time.perf_counter()
    agg_parts: list[pd.DataFrame] = []          # agregados de todos los episodios por trozo
    f84_parts: list[pd.DataFrame] = []          # filas F84 con variables derivadas (sin fechas crudas)
    f84_raw_parts: list[pd.DataFrame] = []      # filas F84 crudas (columnas leídas) para detectar duplicados exactos
    subcode_parts: list[pd.DataFrame] = []
    activity_raw = Counter()
    activity_raw_f84 = Counter()
    sex_raw = Counter()
    n_rows = 0
    n_valid_id_all = 0
    n_candidate_not_exact = 0
    candidate_not_exact_codes = Counter()
    n_dates_unparsed = {"birth": 0, "admission": 0}
    unparsed_examples: Counter = Counter()
    normalisation_checked = False

    for k, chunk in enumerate(C.iter_grd(path, present, chunksize=chunksize)):
        n = len(chunk)
        n_rows += n
        # --- códigos normalizados y profundidad diagnóstica (todos los episodios) ---
        diag = np.empty((n, len(DIAGS)), dtype=object)
        for j, col in enumerate(DIAGS):
            diag[:, j] = normalise_codes(chunk[col])
        if not normalisation_checked:
            sample = chunk[DIAGS].head(2000)
            expected = np.array([[C.normalize_code(v) for v in row] for row in sample.itertuples(index=False)], dtype=object)
            if not np.array_equal(expected, diag[: len(sample)]):
                raise AssertionError("normalise_codes difiere de common.normalize_code")
            normalisation_checked = True
        depth = (diag != "").sum(axis=1).astype(int)
        # --- atributos de todos los episodios ---
        hosp = chunk["COD_HOSPITAL"].fillna("").astype(object).str.strip().to_numpy(dtype=object)
        raw_act = chunk["TIPO_ACTIVIDAD"].fillna("").astype(object).str.strip()
        act_map = {v: classify_activity(v) for v in raw_act.unique()}
        act = raw_act.map(act_map).to_numpy(dtype=object)
        activity_raw.update(raw_act.value_counts().to_dict())
        raw_sex = chunk["SEXO"].fillna("").astype(object).str.strip().str.upper()
        sex_raw.update(raw_sex.value_counts().to_dict())
        sex = raw_sex.map({"HOMBRE": "HOMBRE", "MUJER": "MUJER", "1": "HOMBRE", "2": "MUJER"}).fillna("unknown").to_numpy(dtype=object)
        birth = parse_dates(chunk["FECHA_NACIMIENTO"])
        adm = parse_dates(chunk["FECHA_INGRESO"])
        for name, raw, parsed in (("birth", chunk["FECHA_NACIMIENTO"], birth), ("admission", chunk["FECHA_INGRESO"], adm)):
            bad = parsed.isna().to_numpy()
            n_dates_unparsed[name] += int(bad.sum())
            if bad.any():
                unparsed_examples.update(raw.fillna("<NA>").astype(object)[bad].str.replace(r"\d", "9", regex=True).value_counts().to_dict())
        age = np.floor((adm - birth).dt.days.to_numpy(dtype=float) / 365.25)
        age = np.where((age < 0) | (age > 110), np.nan, age)
        age_group = pd.cut(age, bins=list(range(0, 85, 5)) + [np.inf], right=False, labels=C.AGE_GROUPS).astype(object)
        age_group = np.where(pd.isna(age_group), "unknown", age_group).astype(object)
        ids = chunk[id_col].fillna("").astype(object).str.strip()
        valid = valid_id_mask(ids)
        n_valid_id_all += int(valid.sum())
        agg = (pd.DataFrame({"COD_HOSPITAL": hosp, "activity": act, "sex": sex, "age_group": age_group, "depth": depth})
               .groupby(["COD_HOSPITAL", "activity", "sex", "age_group", "depth"], observed=True).size().rename("n").reset_index())
        agg_parts.append(agg)
        # --- filas candidatas (prefijo F84 en cualquier posición) → pertenencia exacta por variante ---
        cand = np.zeros(n, dtype=bool)
        for j in range(len(DIAGS)):
            cand |= np.char.startswith(diag[:, j].astype(str), "F84")
        if not cand.any():
            continue
        idx = np.flatnonzero(cand)
        sub = diag[idx]
        rec = {"year": year, "COD_HOSPITAL": hosp[idx], "id": ids.to_numpy(dtype=object)[idx], "id_valid": valid[idx],
               "sex": sex[idx], "age": age[idx], "age_group": age_group[idx], "activity": act[idx],
               "activity_raw": raw_act.to_numpy(dtype=object)[idx], "depth": depth[idx],
               "comuna_present": (chunk["COMUNA"].fillna("").astype(object).str.strip() != "").to_numpy(dtype=bool)[idx]
               if "COMUNA" in chunk.columns else np.zeros(len(idx), dtype=bool)}
        any_exact = np.zeros(len(idx), dtype=bool)
        for v, codes in VARIANT_CODES.items():
            principal = np.isin(sub[:, 0], list(codes))
            secondary = np.isin(sub[:, 1:], list(codes)).any(axis=1)
            rec[f"{v}_principal"] = principal
            rec[f"{v}_secondary"] = secondary
            any_exact |= principal | secondary
        f84 = pd.DataFrame(rec)
        # códigos con prefijo F84 que no pertenecen a la lista de config (auditoría; no cuentan como F84)
        prefix_codes = sub[np.char.startswith(sub.astype(str), "F84")]
        for code in prefix_codes:
            if code not in VARIANT_CODES[PRIMARY_VARIANT]:
                candidate_not_exact_codes[code] += 1
        n_candidate_not_exact += int((~any_exact).sum())
        keep = any_exact
        f84_parts.append(f84.loc[keep].reset_index(drop=True))
        f84_raw_parts.append(chunk.iloc[idx[keep]].reset_index(drop=True))
        activity_raw_f84.update(raw_act.iloc[idx[keep]].value_counts().to_dict())
        # subcódigos: un registro cuenta una vez por (subcódigo, posición)
        sub_keep = sub[keep]
        rows = []
        for i in range(sub_keep.shape[0]):
            seen = set()
            for j in range(sub_keep.shape[1]):
                code = sub_keep[i, j]
                if code.startswith("F84"):
                    pos = "principal" if j == 0 else "secondary"
                    seen.add((code, pos))
                    seen.add((code, "any"))
            rows.extend(seen)
        if rows:
            subcode_parts.append(pd.DataFrame(rows, columns=["subcode", "position"]).value_counts().rename("n_episodes").reset_index())
        if max_chunks and k + 1 >= max_chunks:
            print(f"  [prueba] lectura detenida tras {max_chunks} trozo(s)", flush=True)
            break

    agg_all = pd.concat(agg_parts, ignore_index=True).groupby(["COD_HOSPITAL", "activity", "sex", "age_group", "depth"], observed=True)["n"].sum().reset_index()
    f84 = pd.concat(f84_parts, ignore_index=True) if f84_parts else pd.DataFrame()
    f84_raw = pd.concat(f84_raw_parts, ignore_index=True) if f84_raw_parts else pd.DataFrame()
    subcodes = (pd.concat(subcode_parts, ignore_index=True).groupby(["subcode", "position"])["n_episodes"].sum().reset_index()
                if subcode_parts else pd.DataFrame(columns=["subcode", "position", "n_episodes"]))
    n_dup = int(f84_raw.duplicated(keep="first").sum()) if len(f84_raw) else 0
    n_dup_full, dup_differing_columns = 0, []
    if n_dup:
        n_dup_full, dup_differing_columns = verify_full_row_duplicates(path, f84_raw.loc[f84_raw.duplicated(keep=False)], id_col, chunksize)
    elapsed = time.perf_counter() - t0
    stat = path.stat()
    source = {"path": str(path), "bytes": stat.st_size,
              "mtime_utc": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(timespec="seconds"),
              "sha256": C.sha256_file(path) if do_hash else None, "identifier_column": id_col,
              "columns_read": present, "columns_absent": missing}
    print(f"GRD {year}: {n_rows:,} filas; {len(f84):,} episodios con F84 ({PRIMARY_VARIANT}); "
          f"{n_dup} duplicados exactos sobre columnas leídas; {elapsed:,.1f} s", flush=True)
    return dict(year=year, id_col=id_col, agg=agg_all, f84=f84, subcodes=subcodes, n_rows=n_rows,
                n_valid_id_all=n_valid_id_all, n_dup=n_dup, n_dup_full=n_dup_full, dup_differing_columns=dup_differing_columns,
                activity_raw=dict(activity_raw),
                activity_raw_f84=dict(activity_raw_f84), sex_raw=dict(sex_raw),
                n_candidate_not_exact=n_candidate_not_exact, candidate_not_exact_codes=dict(candidate_not_exact_codes),
                dates_unparsed=n_dates_unparsed, unparsed_examples=dict(unparsed_examples), seconds=elapsed, source=source)


# ---------------------------------------------------------------------------
# Derivaciones
# ---------------------------------------------------------------------------
def variant_position_mask(f84: pd.DataFrame, variant: str, position: str) -> np.ndarray:
    p = f84[f"{variant}_principal"].to_numpy(dtype=bool)
    s = f84[f"{variant}_secondary"].to_numpy(dtype=bool)
    return {"any": p | s, "principal": p, "secondary_only": s & ~p, "principal_and_secondary": p & s}[position]


def activity_mask(values: np.ndarray, activity: str) -> np.ndarray:
    return np.ones(len(values), dtype=bool) if activity == "all" else (values == activity)


def build_outputs(years_data: list[dict], names: HospitalNames) -> dict[str, pd.DataFrame]:
    hosp_sets = {d["year"]: set(d["agg"]["COD_HOSPITAL"].unique()) for d in years_data}
    fixed_panel = set.intersection(*[hosp_sets[y] for y in FIXED_PANEL_YEARS if y in hosp_sets]) if all(y in hosp_sets for y in FIXED_PANEL_YEARS) else set()
    summary, hospital, agesex, depthtab, subcode_rows, audit, activity_rows = [], [], [], [], [], [], []
    prev_ids: dict[str, set] = {}
    prev_year: int | None = None
    for d in years_data:
        year, agg, f84 = d["year"], d["agg"], d["f84"]
        hosp_all = agg["COD_HOSPITAL"].to_numpy(dtype=object)
        act_all = agg["activity"].to_numpy(dtype=object)
        n_all = agg["n"].to_numpy(dtype=np.int64)
        depth_all = agg["depth"].to_numpy(dtype=np.int64)
        f_hosp = f84["COD_HOSPITAL"].to_numpy(dtype=object)
        f_act = f84["activity"].to_numpy(dtype=object)
        f_depth = f84["depth"].to_numpy(dtype=np.int64)
        f_id = f84["id"].to_numpy(dtype=object)
        f_valid = f84["id_valid"].to_numpy(dtype=bool)
        panel_masks_all = {"observed": np.ones(len(agg), dtype=bool), "fixed65": np.isin(hosp_all, list(fixed_panel))}
        panel_masks_f84 = {"observed": np.ones(len(f84), dtype=bool), "fixed65": np.isin(f_hosp, list(fixed_panel))}
        # (a) resumen anual --------------------------------------------------------------------
        for panel in PANELS:
            for activity in ACTIVITIES:
                m_all = panel_masks_all[panel] & activity_mask(act_all, activity)
                n_total = int(n_all[m_all].sum())
                depth_mean_all = float((n_all[m_all] * depth_all[m_all]).sum() / n_total) if n_total else float("nan")
                depth_median_all = weighted_median(depth_all[m_all], n_all[m_all]) if n_total else float("nan")
                hospitals_n = int(pd.unique(hosp_all[m_all & (n_all > 0)]).size)
                for variant in VARIANTS:
                    for position in POSITIONS:
                        m = panel_masks_f84[panel] & activity_mask(f_act, activity) & variant_position_mask(f84, variant, position)
                        n_f84 = int(m.sum())
                        ids = f_id[m & f_valid]
                        lo, hi = C.poisson_limits([n_f84])
                        summary.append(dict(
                            year=year, variant=variant, panel=panel, activity=activity, position=position,
                            n_episodes_f84=n_f84, n_episodes_total_same_panel_activity=n_total,
                            rate_per_100k_episodes=1e5 * n_f84 / n_total if n_total else np.nan,
                            rate_lo=1e5 * lo[0] / n_total if n_total else np.nan, rate_hi=1e5 * hi[0] / n_total if n_total else np.nan,
                            persons_within_year=int(pd.unique(ids).size), n_f84_without_valid_id=int((m & ~f_valid).sum()),
                            coding_depth_mean_all=depth_mean_all, coding_depth_median_all=depth_median_all,
                            coding_depth_mean_f84=float(f_depth[m].mean()) if n_f84 else np.nan,
                            coding_depth_median_f84=float(np.median(f_depth[m])) if n_f84 else np.nan,
                            hospitals_n=hospitals_n, identifier_column=d["id_col"]))
        # (b) hospital-año ---------------------------------------------------------------------
        h_tot = agg.groupby("COD_HOSPITAL").apply(lambda g: pd.Series({
            "n_episodes_total": int(g["n"].sum()),
            "n_episodes_hospitalisation": int(g.loc[g.activity == "hospitalisation", "n"].sum()),
            "n_episodes_cma": int(g.loc[g.activity == "cma", "n"].sum()),
            "n_episodes_other": int(g.loc[g.activity == "other", "n"].sum()),
            "coding_depth_mean": float((g["n"] * g["depth"]).sum() / g["n"].sum())}), include_groups=False).reset_index()
        for variant in VARIANTS:
            any_m = variant_position_mask(f84, variant, "any")
            pr_m = variant_position_mask(f84, variant, "principal")
            fv = pd.DataFrame({"COD_HOSPITAL": f_hosp, "any": any_m, "principal": pr_m, "hosp": any_m & (f_act == "hospitalisation"),
                               "cma": any_m & (f_act == "cma"), "depth": np.where(any_m, f_depth, np.nan),
                               "pid": np.where(any_m & f_valid, f_id, None)})
            hv = fv.groupby("COD_HOSPITAL").agg(n_f84_any=("any", "sum"), n_f84_principal=("principal", "sum"),
                                                n_f84_any_hospitalisation=("hosp", "sum"), n_f84_any_cma=("cma", "sum"),
                                                coding_depth_mean_f84=("depth", "mean"),
                                                persons_within_year_f84_any=("pid", lambda s: int(s.dropna().nunique()))).reset_index()
            merged = h_tot.merge(hv, how="left", on="COD_HOSPITAL")
            for col in ["n_f84_any", "n_f84_principal", "n_f84_any_hospitalisation", "n_f84_any_cma", "persons_within_year_f84_any"]:
                merged[col] = merged[col].fillna(0).astype(int)
            merged.insert(0, "variant", variant)
            merged.insert(0, "year", year)
            merged.insert(3, "hospital_name", merged["COD_HOSPITAL"].map(names.name))
            merged.insert(4, "hospital_name_source", merged["COD_HOSPITAL"].map(names.source))
            merged["in_fixed_panel"] = merged["COD_HOSPITAL"].isin(fixed_panel)
            hospital.append(merged)
        # (c) edad × sexo ----------------------------------------------------------------------
        for panel in PANELS:
            for activity in ["all", "hospitalisation"]:
                m_all = panel_masks_all[panel] & activity_mask(act_all, activity)
                den = agg.loc[m_all].groupby(["sex", "age_group"])["n"].sum()
                for variant in VARIANTS:
                    for position in ["any", "principal"]:
                        m = panel_masks_f84[panel] & activity_mask(f_act, activity) & variant_position_mask(f84, variant, position)
                        num = f84.loc[m].groupby(["sex", "age_group"]).size()
                        for sex in SEX_LEVELS:
                            for ag in AGE_LEVELS:
                                agesex.append(dict(year=year, variant=variant, panel=panel, position=position, activity=activity,
                                                   sex=sex, age_group=ag, n_f84=int(num.get((sex, ag), 0)),
                                                   n_total_episodes=int(den.get((sex, ag), 0))))
        # (d) profundidad diagnóstica --------------------------------------------------------
        bins_all = np.array([depth_bin(int(v)) for v in depth_all], dtype=object)
        bins_f84 = np.array([depth_bin(int(v)) for v in f_depth], dtype=object)
        labels = [b[0] for b in DEPTH_BINS]
        for panel in PANELS:
            for activity in ["all", "hospitalisation"]:
                m_all = panel_masks_all[panel] & activity_mask(act_all, activity)
                for variant in VARIANTS:
                    m = panel_masks_f84[panel] & activity_mask(f_act, activity) & variant_position_mask(f84, variant, "any")
                    for label in labels:
                        n_total = int(n_all[m_all & (bins_all == label)].sum())
                        n_f = int((m & (bins_f84 == label)).sum())
                        if n_total == 0 and n_f == 0 and label == "0":
                            continue
                        lo, hi = C.poisson_limits([n_f])
                        depthtab.append(dict(year=year, variant=variant, panel=panel, activity=activity, depth_bin=label,
                                             n_total=n_total, n_f84_any=n_f,
                                             share_f84_within_bin=n_f / n_total if n_total else np.nan,
                                             rate_per_100k_episodes=1e5 * n_f / n_total if n_total else np.nan,
                                             rate_lo=1e5 * lo[0] / n_total if n_total else np.nan,
                                             rate_hi=1e5 * hi[0] / n_total if n_total else np.nan))
        # (e) subcódigos -----------------------------------------------------------------------
        sc = d["subcodes"].copy()
        sc.insert(0, "year", year)
        sc["in_config_f84_subcodes"] = sc["subcode"].isin(CFG.F84_SUBCODES)
        sc["variants"] = sc["subcode"].map(lambda c: "|".join(v for v, codes in VARIANT_CODES.items() if c in codes))
        subcode_rows.append(sc)
        # (f) auditoría de identificadores -----------------------------------------------------
        for variant in VARIANTS:
            m = variant_position_mask(f84, variant, "any")
            ids_valid = pd.Series(f_id[m & f_valid], dtype=object)
            lengths = ids_valid.str.len()
            uniq = set(ids_valid.unique())
            invalid_values = Counter(f_id[m & ~f_valid])
            audit.append(dict(
                year=year, variant=variant, identifier_column=d["id_col"], n_f84=int(m.sum()),
                n_with_valid_id=int((m & f_valid).sum()), n_without_valid_id=int((m & ~f_valid).sum()),
                n_unique_ids=len(uniq),
                id_length_mode=int(lengths.mode().iloc[0]) if len(lengths) else np.nan,
                id_length_min=int(lengths.min()) if len(lengths) else np.nan, id_length_max=int(lengths.max()) if len(lengths) else np.nan,
                invalid_id_values="|".join(f"{k or '<empty>'}:{v}" for k, v in sorted(invalid_values.items())),
                n_f84_ids_shared_with_previous_year=len(uniq & prev_ids.get(variant, set())) if prev_year is not None else np.nan,
                previous_year=prev_year if prev_year is not None else np.nan,
                n_f84_with_comuna=int((m & f84["comuna_present"].to_numpy(dtype=bool)).sum()),
                n_records_total=d["n_rows"], n_records_valid_id_all_episodes=d["n_valid_id_all"],
                n_f84_exact_duplicates_on_read_columns=d["n_dup"]))
            prev_ids[variant] = uniq
        prev_year = year
        # (g) categorías crudas de actividad ---------------------------------------------------
        for raw, n in sorted(d["activity_raw"].items(), key=lambda kv: -kv[1]):
            activity_rows.append(dict(year=year, tipo_actividad_raw=raw if raw else "<empty>", activity=classify_activity(raw),
                                      n_episodes=int(n), n_f84_any_con_rett=int(d["activity_raw_f84"].get(raw, 0))))
    # panel fijo -------------------------------------------------------------------------------
    panel_rows = []
    all_codes = sorted(set.union(*hosp_sets.values()))
    for code in all_codes:
        row = dict(COD_HOSPITAL=code, hospital_name=names.name(code), hospital_name_source=names.source(code), in_fixed_panel=code in fixed_panel)
        for y in sorted(hosp_sets):
            row[f"present_{y}"] = code in hosp_sets[y]
        row["n_years_present"] = sum(code in hosp_sets[y] for y in hosp_sets)
        panel_rows.append(row)
    return {
        "grd_year_summary": pd.DataFrame(summary),
        "grd_hospital_year": pd.concat(hospital, ignore_index=True),
        "grd_age_sex_year": pd.DataFrame(agesex),
        "grd_coding_depth_year": pd.DataFrame(depthtab),
        "grd_subcode_year": pd.concat(subcode_rows, ignore_index=True).sort_values(["year", "subcode", "position"]).reset_index(drop=True),
        "grd_identifier_audit": pd.DataFrame(audit),
        "grd_activity_categories_year": pd.DataFrame(activity_rows),
        "grd_fixed_panel_hospitals": pd.DataFrame(panel_rows),
        "_fixed_panel": fixed_panel, "_hosp_sets": hosp_sets,
    }


# ---------------------------------------------------------------------------
# Controles
# ---------------------------------------------------------------------------
def control_row(name: str, key, expected, observed, note: str = "") -> dict:
    row = dict(name=name, key=key, expected=expected, observed=observed, abs_diff=np.nan, rel_diff=np.nan, status="missing", note=note)
    if observed is None or (isinstance(observed, float) and np.isnan(observed)):
        return row
    if isinstance(expected, (bool, np.bool_)) or isinstance(observed, (bool, np.bool_)):
        row["status"] = "ok" if bool(expected) == bool(observed) else "differs"
        return row
    diff = float(observed) - float(expected)
    row["abs_diff"] = diff
    row["rel_diff"] = diff / float(expected) if float(expected) != 0 else np.nan
    if isinstance(expected, (int, np.integer)):
        row["status"] = "ok" if diff == 0 else "differs"
    else:
        decimals = len(str(expected).split(".")[1]) if "." in str(expected) else 0
        rounded_equal = round(float(observed), decimals) == round(float(expected), decimals)
        row["status"] = "ok" if rounded_equal or abs(row["rel_diff"]) <= 0.005 else "differs"
    return row


def build_controls(out: dict[str, pd.DataFrame], years_data: list[dict]) -> pd.DataFrame:
    s = out["grd_year_summary"]
    aud = out["grd_identifier_audit"]

    def pick(year, panel="observed", activity="all", position="any", variant=PRIMARY_VARIANT, col="n_episodes_f84"):
        r = s.loc[(s.year == year) & (s.panel == panel) & (s.activity == activity) & (s.position == position) & (s.variant == variant), col]
        return None if r.empty else r.iloc[0]

    rows = []
    ctl = CFG.CONTROLS
    nrows = {d["year"]: d["n_rows"] for d in years_data}
    for y, exp in ctl["grd_records_total"].items():
        rows.append(control_row("grd_records_total", y, exp, nrows.get(y), "filas de datos leídas (sin eliminar duplicados)"))
    for y, exp in ctl["grd_hospitals_observed"].items():
        rows.append(control_row("grd_hospitals_observed", y, exp, pick(y, col="hospitals_n"), "COD_HOSPITAL distintos con ≥1 episodio"))
    for y, exp in ctl["grd_f84_any"].items():
        rows.append(control_row("grd_f84_any", y, exp, None if pick(y) is None else int(pick(y)), f"variant={PRIMARY_VARIANT}; observed; all; any; sin eliminar duplicados"))
    for y, exp in ctl["grd_f84_any_panel65"].items():
        rows.append(control_row("grd_f84_any_panel65", y, exp, None if pick(y, panel="fixed65") is None else int(pick(y, panel="fixed65")), "panel fijo = hospitales presentes en 2019–2022"))
    for y, exp in ctl["grd_f84_any_strict_hospitalisation"].items():
        v = pick(y, activity="hospitalisation")
        v_fix = pick(y, activity="hospitalisation", panel="fixed65")
        v_cma = pick(y, activity="cma")
        rows.append(control_row("grd_f84_any_strict_hospitalisation", y, exp, None if v is None else int(v),
                                "panel observado; TIPO_ACTIVIDAD == HOSPITALIZACIÓN. Desde 2020 solo existen dos categorías, por lo que observado = any − CMA "
                                f"({pick(y)} − {v_cma}); el valor esperado del brief en 2023–2024 corresponde al panel fijo de 65 (fila _fixed65)"))
        rows.append(control_row("grd_f84_any_strict_hospitalisation_fixed65", y, exp, None if v_fix is None else int(v_fix),
                                "panel fijo de 65; TIPO_ACTIVIDAD == HOSPITALIZACIÓN; reproduce los valores del brief en todos los años"))
    for y, exp in ctl["grd_cma"].items():
        v = pick(y, activity="cma")
        rows.append(control_row("grd_cma", y, exp, None if v is None else int(v), "TIPO_ACTIVIDAD == CIRUGÍA MAYOR AMBULATORIA (CMA)"))
    for y, exp in ctl["grd_f84_principal"].items():
        v = pick(y, position="principal")
        rows.append(control_row("grd_f84_principal", y, exp, None if v is None else int(v), "DIAGNOSTICO1 ∈ F84 (con o sin F84 secundario)"))
    any24, sec24 = pick(2024), pick(2024, position="secondary_only")
    rows.append(control_row("grd_f84_secondary_only_share_2024", 2024, ctl["grd_f84_secondary_only_share_2024"],
                            None if (any24 is None or not any24) else float(sec24) / float(any24),
                            f"secondary_only/any = {sec24}/{any24}" if any24 else ""))
    for y, exp in ctl["grd_coding_depth_all_mean"].items():
        rows.append(control_row("grd_coding_depth_all_mean", y, exp, pick(y, col="coding_depth_mean_all"), "media de DIAGNOSTICO no vacíos, todos los episodios (panel observado)"))
    for y, exp in ctl["grd_coding_depth_f84_mean"].items():
        rows.append(control_row("grd_coding_depth_f84_mean", y, exp, pick(y, col="coding_depth_mean_f84"), f"media entre episodios F84 ({PRIMARY_VARIANT}, any)"))
    for y, exp in ctl["grd_persons_within_year_f84_any"].items():
        n_u = pick(y, col="persons_within_year")
        n_inv = pick(y, col="n_f84_without_valid_id")
        inv_vals = aud.loc[(aud.year == y) & (aud.variant == PRIMARY_VARIANT), "invalid_id_values"]
        inv_vals = "" if inv_vals.empty or pd.isna(inv_vals.iloc[0]) else str(inv_vals.iloc[0])
        rows.append(control_row("grd_persons_within_year_f84_any", y, exp, None if n_u is None else int(n_u),
                                f"identificadores válidos distintos dentro del año; {0 if n_inv is None else int(n_inv)} episodios F84 sin identificador válido excluidos"
                                + (f" ({inv_vals})" if inv_vals else "") + "; nunca se deduplica entre años"))
        if n_u is not None:
            rows.append(control_row("grd_persons_within_year_f84_any_placeholder_as_one_id", y, exp, int(n_u) + (1 if int(n_inv) > 0 else 0),
                                    "diagnóstico: el valor del brief cuenta el marcador de identificador inválido como una persona adicional; no se usa en las salidas"))
    fixed = out["_fixed_panel"]
    hs = out["_hosp_sets"]
    rows.append(control_row("grd_fixed_panel_size", "2019-2022", 65, len(fixed), "intersección de COD_HOSPITAL presentes en 2019, 2020, 2021 y 2022"))
    for y in (2023, 2024):
        if y in hs:
            rows.append(control_row("grd_fixed_panel_subset_of_year", y, True, fixed.issubset(hs[y]), f"{len(fixed - hs[y])} hospitales del panel ausentes en {y}"))
    for d in years_data:
        y = d["year"]
        n_before = int(aud.loc[(aud.year == y) & (aud.variant == PRIMARY_VARIANT), "n_f84"].iloc[0])
        rows.append(control_row("grd_f84_exact_duplicates_on_read_columns", y, 0, d["n_dup"],
                                f"duplicados sobre las {len(d['source']['columns_read'])} columnas leídas: episodios F84 antes={n_before}, después={n_before - d['n_dup']}; "
                                "las salidas y los controles usan 'antes' (sin eliminación)"
                                + (f"; las filas difieren en {', '.join(d['dup_differing_columns'])}" if d["dup_differing_columns"] else "")))
        rows.append(control_row("grd_f84_exact_duplicates_full_row", y, 0, d["n_dup_full"],
                                "duplicados exactos sobre las 129 columnas (verificación dirigida solo si hay duplicados sobre columnas leídas); "
                                "expected=0 coincide con output_files/consolidacion/grd_epi_quality.csv"))
        rows.append(control_row("grd_f84_prefix_codes_outside_config_list", y, 0, d["n_candidate_not_exact"],
                                "filas con algún código de prefijo F84 y ninguna pertenencia exacta a config.F84_SUBCODES: "
                                + (", ".join(f"{k}:{v}" for k, v in sorted(d["candidate_not_exact_codes"].items())) or "ninguno")))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Principal
# ---------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--years", type=int, nargs="+", default=CFG.YEARS_GRD)
    ap.add_argument("--chunksize", type=int, default=200_000)
    ap.add_argument("--max-chunks", type=int, default=None, help="solo pruebas: limita los trozos leídos por año")
    ap.add_argument("--no-hash", action="store_true", help="omite el SHA-256 de los archivos fuente")
    args = ap.parse_args(argv)

    t_start = time.perf_counter()
    started = datetime.now(tz=timezone.utc).isoformat(timespec="seconds")
    print(f"[{MODULE}] GRD en {CFG.PATHS['grd']}; años {args.years}; trozos de {args.chunksize:,} filas", flush=True)
    names = HospitalNames()
    print(f"  catálogo de hospitales GRD: {len(names.grd)} códigos ({names.grd_path or 'no disponible'}); "
          f"respaldo DEIS: {len(names.deis)} códigos ({names.deis_path or 'no disponible'})", flush=True)
    years_data = [read_year(y, args.chunksize, args.max_chunks, not args.no_hash) for y in args.years]
    out = build_outputs(years_data, names)
    controls = build_controls(out, years_data)

    written = {}
    for name, df in out.items():
        if name.startswith("_"):
            continue
        path = C.atomic_write_csv(df, CFG.TIDY / f"{name}.csv")
        written[name] = {"path": str(path), "rows": int(len(df)), "columns": list(df.columns)}
    ctl_path = C.atomic_write_csv(controls, CONTROLS_DIR / f"{MODULE}_controls.csv")
    written["controls"] = {"path": str(ctl_path), "rows": int(len(controls)), "columns": list(controls.columns)}

    runlog = {
        "module": MODULE, "script": SCRIPT, "started_utc": started, "seconds_total": round(time.perf_counter() - t_start, 1),
        "args": vars(args), "hospital_catalogue": names.grd_path, "hospital_catalogue_fallback": names.deis_path,
        "hospital_name_sources": {c: names.source(c) for c in sorted(set.union(*out["_hosp_sets"].values())) if names.source(c) != "grd_master"},
        "fixed_panel": sorted(out["_fixed_panel"]), "hospitals_per_year": {str(y): len(s) for y, s in out["_hosp_sets"].items()},
        "years": [{k: v for k, v in d.items() if k not in {"agg", "f84", "subcodes"}} for d in years_data],
        "outputs": written, "controls_status": controls.status.value_counts().to_dict(),
    }
    C.atomic_write_json(runlog, CONTROLS_DIR / f"{MODULE}_runlog.json")
    print(controls[["name", "key", "expected", "observed", "status"]].to_string(index=False), flush=True)
    print(f"[{MODULE}] listo en {runlog['seconds_total']:,.1f} s; controles: {runlog['controls_status']}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())

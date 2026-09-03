"""Extract a year-specific REM catalogue and aggregate the selected records."""
from __future__ import annotations

import argparse
from pathlib import Path
import re
import pandas as pd
from openpyxl import load_workbook
from audit_grd_linkage import data_root

A05_CODES = {"06902600", "05225000"} | {f"059900{i}" for i in range(22, 32)}
AGE_GROUPS = [f"{i}-{i+4}" for i in range(0, 80, 5)] + ["80+"]


def catalogue(root: Path, years: list[int]) -> pd.DataFrame:
    records = []
    for year in years:
        paths = [p for p in (root / "REM/SerieA/metadata/diccionarios" / str(year)).glob("*")
                 if p.name.startswith(("SA", "DICCIONARIO CODIGOS SA"))]
        if len(paths) != 1:
            raise ValueError(f"Diccionario Serie A no unívoco para {year}: {paths}")
        path = paths[0]
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in ["A03", "A05", "A28"]:
            if sheet not in workbook.sheetnames:
                continue
            section, headers, first_data = "", [], False
            for number, row in enumerate(workbook[sheet].iter_rows(max_col=60, values_only=True), 1):
                text = " | ".join(str(v).strip().replace("\n", " ") for v in row[:4] if v is not None)
                if "SECCIÓN" in text.upper():
                    section, headers, first_data = text, [], False
                code = str(row[0]).strip().zfill(8)
                is_data = bool(re.fullmatch(r"\d{8}", code)) and "COL01" in row
                if not is_data:
                    if not first_data:
                        headers.append(row)
                    continue
                first_data = True
                selected = (sheet == "A05" and code in A05_CODES) or (
                    sheet == "A03" and "TAMIZAJE TRASTORNO ESPECTRO AUTISTA" in section.upper()
                ) or (sheet == "A28" and "AUTISTA" in text.upper())
                if not selected:
                    continue
                label = " | ".join(str(v).strip().replace("\n", " ") for v in row[1:4]
                                   if v is not None and not str(v).startswith("COL"))
                if sheet == "A05":
                    # Validate the actual column tokens against their section headers.
                    p2, p3, p4 = [row.index(c) for c in ["COL02", "COL03", "COL04"]]
                    assert any(str(h[p2]).strip().lower() == "hombres" for h in headers), (year, code)
                    assert any(str(h[p3]).strip().lower() == "mujeres" for h in headers), (year, code)
                    assert any(re.search(r"0\s*-\s*4", str(h[p4])) for h in headers), (year, code)
                    for group_index, age in enumerate(AGE_GROUPS):
                        pm, pf = [row.index(f"COL{4 + 2*group_index + s:02d}") for s in [0, 1]]
                        assert any(str(h[pm]).strip().lower() == "hombres" for h in headers), (year, code, age)
                        assert any(str(h[pf]).strip().lower() == "mujeres" for h in headers), (year, code, age)
                        expected = rf"{5*group_index}\s*(?:-|a)\s*{5*group_index+4}" if group_index < 16 else r"80\s*y"
                        assert any(re.search(expected, str(h[pm]).lower()) for h in headers), (year, code, age)
                elif sheet == "A03":
                    p1, p2 = [row.index(c) for c in ["COL01", "COL02"]]
                    assert any(str(h[p1]).strip().lower() == "hombres" for h in headers), (year, code)
                    assert any(str(h[p2]).strip().lower() == "mujeres" for h in headers), (year, code)
                else:
                    p1 = row.index("COL01")
                    assert any(str(h[p1]).strip().lower() == "ambos sexos" for h in headers), (year, code)
                records.append(dict(year=year, sheet=sheet, code=code, label=label, section=section,
                                    dictionary=str(path.relative_to(root)), excel_row=number,
                                    total_rule="Col01+Col02" if sheet == "A03" else "Col01",
                                    age_sex_mapping="A05_Col04_Col37" if sheet == "A05" else "not_applied"))
        workbook.close()
    return pd.DataFrame(records).drop_duplicates(["year", "sheet", "code"])


def run(root: Path, output: Path, years: list[int]) -> None:
    output.mkdir(parents=True, exist_ok=True)
    cat = catalogue(root, years)
    cat.to_csv(output / "rem_code_catalogue.csv", index=False)
    numeric = [f"Col{i:02d}" for i in range(1, 51)]
    keys = ["Ano", "Mes", "IdServicio", "IdEstablecimiento", "CodigoPrestacion", "IdRegion", "IdComuna"]
    all_totals, annual, ages, audits = [], [], [], []
    region_ages, comunas, panels = [], [], []
    for year in years:
        path = root / "REM/SerieA" / f"SerieA_{year}.csv"
        codes = set(cat.loc[cat.year == year, "code"])
        selected, nrows, bad_year, bad_month = [], 0, 0, 0
        for chunk in pd.read_csv(path, sep=";", dtype="string", chunksize=200000):
            nrows += len(chunk)
            bad_year += int((pd.to_numeric(chunk.Ano, errors="coerce") != year).sum())
            months = pd.to_numeric(chunk.Mes, errors="coerce")
            bad_month += int((~months.between(1, 12)).sum())
            chunk["CodigoPrestacion"] = chunk.CodigoPrestacion.str.strip().str.zfill(8)
            selected.append(chunk.loc[chunk.CodigoPrestacion.isin(codes)].copy())
        data = pd.concat(selected, ignore_index=True)
        duplicate_rows = int(data.duplicated().sum())
        data = data.drop_duplicates()
        conflicts = int(data.duplicated(keys, keep=False).sum())
        if conflicts or bad_year or bad_month:
            raise ValueError(f"REM {year}: {conflicts} claves duplicadas, {bad_year} años y {bad_month} meses inválidos")
        missing = data[numeric].isna() | data[numeric].isin(["", "-", ".", "NA"])
        values = data[numeric].mask(missing).apply(lambda x: pd.to_numeric(x.str.replace(",", ".", regex=False), errors="coerce"))
        invalid_numeric = int((values.isna() & ~missing).sum().sum())
        if invalid_numeric or (values < 0).any().any():
            raise ValueError(f"REM {year}: valores inválidos o negativos; revisar antes de agregar")
        data[numeric] = values
        a03_codes = set(cat.loc[(cat.year == year) & (cat.sheet == "A03"), "code"])
        data["total_known"] = data.Col01
        is_a03 = data.CodigoPrestacion.isin(a03_codes)
        data.loc[is_a03, "total_known"] = data.loc[is_a03, ["Col01", "Col02"]].sum(axis=1, min_count=2)
        measures = numeric + ["total_known"]
        data["year"] = year
        data["Mes"] = pd.to_numeric(data.Mes)
        summary = data.groupby(["year", "Mes", "IdRegion", "CodigoPrestacion"], dropna=False)[measures].sum(min_count=1).reset_index()
        coverage = data.groupby(["year", "Mes", "IdRegion", "CodigoPrestacion"], dropna=False).agg(
            reporting_establishments=("IdEstablecimiento", "nunique"), rows=("Col01", "size"),
            rows_with_Col01=("Col01", "count"), rows_total_known=("total_known", "count")).reset_index()
        summary = summary.merge(coverage, validate="one_to_one")
        all_totals.append(summary)
        grouped = data.groupby("CodigoPrestacion")[measures].sum(min_count=1).reset_index().rename(columns={"CodigoPrestacion": "code"})
        cov = data.groupby("CodigoPrestacion").agg(months=("Mes", "nunique"),
                    reporting_establishments=("IdEstablecimiento", "nunique"), rows=("Col01", "size"),
                    rows_total_known=("total_known", "count")).reset_index().rename(columns={"CodigoPrestacion": "code"})
        joined = cat.loc[cat.year == year].merge(grouped, how="left", on="code", validate="one_to_one").merge(cov, how="left", on="code")
        annual.append(joined)
        a05 = joined.loc[joined.sheet == "A05"]
        for row in a05.to_dict("records"):
            for index, age in enumerate(AGE_GROUPS):
                for sex_index, sex in enumerate(["Hombres", "Mujeres"]):
                    ages.append(dict(year=year, code=row["code"], label=row["label"], section=row["section"],
                                     age_group=age, sex=sex, count=row[f"Col{4 + 2*index + sex_index:02d}"]))
        a05_data = data.loc[data.CodigoPrestacion.isin(set(a05.code))]
        # Regional age/sex sums of the A05 columns 04–37 (observed cells only).
        region_cells = a05_data.groupby(["year", "CodigoPrestacion", "IdRegion"], dropna=False)[numeric[3:37]].sum(min_count=1).reset_index()
        for row in region_cells.to_dict("records"):
            for index, age in enumerate(AGE_GROUPS):
                for sex_index, sex in enumerate(["Hombres", "Mujeres"]):
                    region_ages.append(dict(year=year, code=row["CodigoPrestacion"], IdRegion=row["IdRegion"],
                                            age_group=age, sex=sex, count=row[f"Col{4 + 2*index + sex_index:02d}"]))
        # Comuna of the reporting establishment: sums and coverage, all selected codes.
        comuna = data.groupby(["year", "CodigoPrestacion", "IdRegion", "IdComuna"], dropna=False).agg(
            Col01=("Col01", lambda s: s.sum(min_count=1)), Col02=("Col02", lambda s: s.sum(min_count=1)),
            Col03=("Col03", lambda s: s.sum(min_count=1)), total_known=("total_known", lambda s: s.sum(min_count=1)),
            rows=("Col01", "size"), rows_total_known=("total_known", "count"),
            reporting_establishments=("IdEstablecimiento", "nunique")).reset_index().rename(columns={"CodigoPrestacion": "code"})
        comunas.append(comuna)
        # Establishment panel: months with a row, months with a positive total, sums by establishment and code.
        panel = data.assign(positive=data.total_known > 0).groupby(
            ["year", "CodigoPrestacion", "IdServicio", "IdRegion", "IdComuna", "IdEstablecimiento"], dropna=False).agg(
            months_with_rows=("Mes", "nunique"), rows=("Col01", "size"), rows_total_known=("total_known", "count"),
            months_positive=("positive", "sum"), total_known=("total_known", lambda s: s.sum(min_count=1)),
            Col01=("Col01", lambda s: s.sum(min_count=1))).reset_index().rename(columns={"CodigoPrestacion": "code"})
        panels.append(panel)
        valid_sex = a05_data[["Col01", "Col02", "Col03"]].notna().all(axis=1)
        valid_age = a05_data[["Col01"] + numeric[3:37]].notna().all(axis=1)
        audits.append(dict(year=year, records=nrows, selected_records=len(data),
                           exact_duplicates_removed=duplicate_rows, conflicting_key_rows=conflicts,
                           expected_codes=len(codes), observed_codes=data.CodigoPrestacion.nunique(),
                           a05_rows_sex_comparable=int(valid_sex.sum()),
                           a05_rows_sex_mismatch=int((a05_data.loc[valid_sex,"Col01"] != a05_data.loc[valid_sex,["Col02","Col03"]].sum(axis=1)).sum()),
                           a05_rows_age_comparable=int(valid_age.sum()),
                           a05_rows_age_mismatch=int((a05_data.loc[valid_age,"Col01"] != a05_data.loc[valid_age,numeric[3:37]].sum(axis=1)).sum()),
                           bytes=path.stat().st_size, mtime_ns=path.stat().st_mtime_ns))
        print(f"REM {year}: {nrows:,} filas; {len(data):,} seleccionadas", flush=True)
    pd.concat(all_totals).to_csv(output / "rem_monthly_region_raw_columns.csv", index=False)
    pd.concat(annual).to_csv(output / "rem_annual_by_code.csv", index=False)
    pd.DataFrame(ages).to_csv(output / "rem_a05_age_sex.csv", index=False)
    pd.DataFrame(audits).to_csv(output / "rem_quality.csv", index=False)
    pd.DataFrame(region_ages).to_csv(output / "rem_a05_region_age_sex.csv", index=False)
    pd.concat(comunas).to_csv(output / "rem_annual_comuna.csv", index=False)
    pd.concat(panels).to_csv(output / "rem_establishment_panel.csv", index=False)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("output_files/consolidacion"))
    parser.add_argument("--years", type=int, nargs="+", default=list(range(2017, 2025)))
    args = parser.parse_args()
    run(data_root(), args.output, args.years)

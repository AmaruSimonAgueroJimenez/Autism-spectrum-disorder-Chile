"""Descriptive epidemiology aggregates for hospital discharges with F84 codes (GRD 2019–2024).

Reads each canonical GRD file once, in chunks, keeps only the rows with a
F84 code in any of the 35 diagnostic positions and writes **aggregates only**:
no identifiers, no dates and no record-level rows leave this script.

Units are kept explicit in every table:
- `records`: discharge records after exact deduplication (one hospitalisation
  episode as published by the producer);
- `identifiers`/`persons`: distinct valid identifiers within the cell, never
  summed across cells or across the 2019–2020 and 2021–2024 identifier eras.

The three case definitions of `grd_trajectories.py` are reused so that every
document shares one vocabulary: `autismo_F840`, `TEA_operacional` and
`F84_historico`. Positions follow the same operational rule: DIAGNOSTICO1 is
treated as principal and DIAGNOSTICO2–35 as secondary.
"""
from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd

from audit_grd_linkage import data_root, parse_dates, valid_id
from grd_trajectories import DIAGS, DEFINITIONS, ERAS, normalize_icd

AGE_GROUPS = [f"{i}-{i + 4}" for i in range(0, 80, 5)] + ["80+"]
AGE_BANDS = ["0-4", "5-9", "10-14", "15-19", "20-29", "30-44", "45+"]
LOS_BINS = ["0", "1", "2", "3-4", "5-7", "8-14", "15-30", "31-90", "91+"]
ROLES = ["principal", "solo_secundario", "principal_y_secundario"]
SUBCODES = ["F84", "F840", "F841", "F842", "F843", "F844", "F845", "F848", "F849"]
FEATURES = ["TIPO_INGRESO", "TIPO_ACTIVIDAD", "PREVISION", "TIPOALTA", "ESPECIALIDAD_MEDICA",
            "TIPO_PROCEDENCIA", "SERVICIO_SALUD", "NACIONALIDAD", "ETNIA",
            "IR_29301_SEVERIDAD", "IR_29301_MORTALIDAD"]
USECOLS = ["COD_HOSPITAL", "SEXO", "FECHA_NACIMIENTO", "PROVINCIA", "COMUNA", "FECHA_INGRESO",
           "FECHAALTA", "IR_29301_PESO"] + FEATURES + DIAGS


def strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def normalize_name(value: object) -> str:
    """Upper-case ASCII name without punctuation; keeps a stable key for comunas."""
    if pd.isna(value):
        return ""
    text = strip_accents(str(value)).upper().replace("'", " ").replace("-", " ")
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", text)).strip()


def age_group(age: pd.Series) -> pd.Series:
    bins = list(range(0, 85, 5)) + [np.inf]
    return pd.cut(age, bins=bins, right=False, labels=AGE_GROUPS)


def age_band(age: pd.Series) -> pd.Series:
    return pd.cut(age, bins=[0, 5, 10, 15, 20, 30, 45, np.inf], right=False, labels=AGE_BANDS)


def los_bin(days: pd.Series) -> pd.Series:
    return pd.cut(days, bins=[-0.5, 0.5, 1.5, 2.5, 4.5, 7.5, 14.5, 30.5, 90.5, np.inf], labels=LOS_BINS)


def group_prevision(value: object) -> str:
    text = "" if pd.isna(value) else str(value).upper()
    if "FONASA" in text and "MAI" in text:
        for letter in "ABCD":
            if text.rstrip().endswith(f" {letter}"):
                return f"FONASA MAI {letter}"
        return "FONASA MAI (tramo no informado)"
    if "FONASA" in text:
        return "FONASA libre elección"
    if "ISAPRE" in text:
        return "ISAPRE"
    if any(k in text for k in ["DIPRECA", "CAPREDENA", "FFAA", "SISA"]):
        return "FFAA y de Orden"
    if "PARTICULAR" in text:
        return "Particular"
    if text.strip() in {"", "NO IDENTIFICADA", "DESCONOCIDO", "NAN"}:
        return "No identificada"
    return "Otra"


def group_nationality(value: object) -> str:
    text = "" if pd.isna(value) else str(value).upper().strip()
    if text == "CHILE":
        return "Chile"
    if text in {"", "DESCONOCIDO", "DESCONOCIDA", "NAN"}:
        return "Desconocida"
    return "Otro país"


def group_ethnicity(value: object) -> str:
    text = "" if pd.isna(value) else strip_accents(str(value)).upper().strip()
    if text in {"NINGUNO", "NINGUNA", "OTRO", "OTRA", "OTROS"}:
        return "Ninguno/otro"
    if text in {"", "DESCONOCIDO", "NAN", "NO INFORMADO"}:
        return "Desconocido"
    return "Pueblo originario declarado"


def read_year(root: Path, year: int, chunksize: int, max_chunks: int | None) -> tuple[pd.DataFrame, dict]:
    path = root / "GRD" / f"GRD_PUBLICO_{year}.csv"
    header = pd.read_csv(path, sep="|", nrows=0).columns
    id_col = "ID_BENEFICIARIO" if "ID_BENEFICIARIO" in header else "CIP_ENCRIPTADO"
    selected, scanned = [], 0
    reader = pd.read_csv(path, sep="|", dtype="string", usecols=[id_col] + USECOLS, chunksize=chunksize)
    for number, chunk in enumerate(reader):
        scanned += len(chunk)
        mask = pd.Series(False, index=chunk.index)
        for column in DIAGS:
            mask |= chunk[column].str.match(r"(?i)^\s*F\.?84", na=False)
        selected.append(chunk.loc[mask].rename(columns={id_col: "id"}))
        if max_chunks and number + 1 >= max_chunks:
            break
    frame = pd.concat(selected, ignore_index=True)
    prefiltered = len(frame)
    for column in DIAGS:
        frame[column] = frame[column].map(normalize_icd)
    exact = frame[DIAGS].isin(DEFINITIONS["F84_historico"]).any(axis=1)
    frame = frame.loc[exact].copy()
    duplicates = int(frame.duplicated().sum())
    frame = frame.drop_duplicates().reset_index(drop=True)
    frame["year"] = year
    frame["id"] = frame["id"].str.strip()
    frame["valid_id"] = valid_id(frame["id"])
    frame.loc[~frame.valid_id, "id"] = pd.NA
    frame["sex"] = frame.SEXO.str.strip().str.upper().replace({"1": "HOMBRE", "2": "MUJER"})
    frame.loc[~frame.sex.isin(["HOMBRE", "MUJER"]), "sex"] = pd.NA
    frame["birth"] = parse_dates(frame.FECHA_NACIMIENTO)
    frame["admission"] = parse_dates(frame.FECHA_INGRESO)
    frame["discharge"] = parse_dates(frame.FECHAALTA)
    frame["valid_dates"] = frame.admission.notna() & frame.discharge.notna() & (frame.discharge >= frame.admission)
    age_days = (frame.admission - frame.birth).dt.days
    frame["age"] = np.floor(age_days / 365.25)
    implausible = frame.age.notna() & ((frame.age < 0) | (frame.age > 110))
    frame.loc[implausible, "age"] = np.nan
    frame["age_group"] = age_group(frame.age).astype("object")
    frame["age_band"] = age_band(frame.age).astype("object")
    frame["los"] = (frame.discharge - frame.admission).dt.days.where(frame.valid_dates)
    frame["los_bin"] = los_bin(frame.los).astype("object")
    frame["month"] = frame.admission.dt.month
    frame["comuna_norm"] = frame.COMUNA.map(normalize_name)
    frame["provincia_norm"] = frame.PROVINCIA.map(normalize_name)
    frame["prevision_grupo"] = frame.PREVISION.map(group_prevision)
    frame["nacionalidad_grupo"] = frame.NACIONALIDAD.map(group_nationality)
    frame["etnia_grupo"] = frame.ETNIA.map(group_ethnicity)
    frame["fallecido"] = frame.TIPOALTA.str.upper().str.contains("FALLECIDO", na=False)
    frame["peso_grd"] = pd.to_numeric(frame.IR_29301_PESO.str.replace(",", ".", regex=False), errors="coerce")
    for name, codes in DEFINITIONS.items():
        principal = frame.DIAGNOSTICO1.isin(codes)
        secondary = frame[DIAGS[1:]].isin(codes).any(axis=1)
        role = pd.Series("sin_codigo", index=frame.index, dtype="object")
        role[principal & secondary] = "principal_y_secundario"
        role[principal & ~secondary] = "principal"
        role[~principal & secondary] = "solo_secundario"
        frame[f"role_{name}"] = role
    quality = dict(year=year, identifier_column=id_col, records_scanned=scanned, prefiltered_rows=prefiltered,
                   f84_records=int(exact.sum()), exact_duplicates_removed=duplicates,
                   records_kept=len(frame), records_missing_id=int((~frame.valid_id).sum()),
                   records_missing_sex=int(frame.sex.isna().sum()),
                   records_invalid_dates=int((~frame.valid_dates).sum()),
                   records_missing_birth=int(frame.birth.isna().sum()),
                   records_implausible_age=int(implausible.sum()),
                   records_missing_age_group=int(frame.age_group.isna().sum()),
                   records_missing_comuna=int((frame.comuna_norm == "").sum()),
                   records_deaths=int(frame.fallecido.sum()))
    print(f"GRD {year}: {scanned:,} filas leídas; {len(frame):,} registros F84", flush=True)
    return frame, quality


def count(frame: pd.DataFrame, keys: list[str], name: str = "records") -> pd.DataFrame:
    return frame.groupby(keys, dropna=False, observed=True).size().reset_index(name=name)


def with_definition(frame: pd.DataFrame, name: str) -> pd.DataFrame:
    """Rows carrying a code of the definition, with their role in that definition."""
    subset = frame.loc[frame[f"role_{name}"] != "sin_codigo"].copy()
    subset["definition"] = name
    subset["role"] = subset[f"role_{name}"]
    return subset


def hospital_catalogue(root: Path) -> pd.DataFrame:
    path = root / "GRD" / "metadata" / "códigos grd.xlsx"
    if not path.is_file():
        return pd.DataFrame(columns=["COD_HOSPITAL", "hospital"])
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for row in workbook["Hospitales"].iter_rows(max_col=2, values_only=True):
        if row[0] is not None and re.fullmatch(r"\d{6}", str(row[0]).strip()) and row[1]:
            rows.append(dict(COD_HOSPITAL=str(row[0]).strip(), hospital=str(row[1]).strip()))
    workbook.close()
    return pd.DataFrame(rows).drop_duplicates("COD_HOSPITAL")


def run(root: Path, output: Path, years: list[int], chunksize: int, max_chunks: int | None) -> None:
    output.mkdir(parents=True, exist_ok=True)
    frames, quality = [], []
    for year in years:
        frame, qc = read_year(root, year, chunksize, max_chunks)
        frames.append(frame)
        quality.append(qc)
    data = pd.concat(frames, ignore_index=True)
    data["era"] = data.year.map({year: era for era, era_years in ERAS.items() for year in era_years})
    tables: dict[str, list[pd.DataFrame]] = {}

    def add(name: str, frame: pd.DataFrame) -> None:
        tables.setdefault(name, []).append(frame)

    for name in DEFINITIONS:
        subset = with_definition(data, name)
        keys = ["year", "definition", "role", "sex", "age_group", "comuna_norm"]
        strata = subset.groupby(keys, dropna=False, observed=True).agg(
            records=("year", "size"), identifiers=("id", "nunique"),
            records_missing_id=("valid_id", lambda s: int((~s).sum()))).reset_index()
        add("strata", strata)
        add("age_single", count(subset, ["year", "definition", "role", "sex", "age"]))
        add("monthly", count(subset, ["year", "definition", "role", "month"]))
        add("hospitals", subset.groupby(["year", "definition", "role", "COD_HOSPITAL"], dropna=False).agg(
            records=("year", "size"), identifiers=("id", "nunique")).reset_index())
        # Discharge characteristics: raw producer categories plus grouped versions.
        for variable in FEATURES + ["prevision_grupo", "nacionalidad_grupo", "etnia_grupo", "age_band", "sex", "los_bin"]:
            values = subset[variable].astype("object").where(subset[variable].notna(), "No informado")
            counted = subset.assign(value=values.astype(str).str.strip()).groupby(
                ["year", "definition", "role", "value"], dropna=False).size().reset_index(name="records")
            counted.insert(3, "variable", variable)
            add("features", counted)
        stays = subset.loc[subset.los.notna()]
        los_stats = dict(n="size", mean="mean", sd="std", median="median", q25=lambda s: s.quantile(.25),
                         q75=lambda s: s.quantile(.75), p90=lambda s: s.quantile(.9), max="max", total_days="sum",
                         n_los0=lambda s: int((s == 0).sum()))
        add("los", stays.groupby(["year", "definition", "role", "TIPO_ACTIVIDAD", "age_band"], dropna=False, observed=True).los.agg(**los_stats).reset_index())
        add("los_summary", stays.groupby(["year", "definition", "role", "TIPO_ACTIVIDAD"], dropna=False).los.agg(**los_stats).reset_index())
        add("los_age", stays.groupby(["definition", "role", "TIPO_ACTIVIDAD", "age_band"], dropna=False, observed=True).los.agg(**los_stats).reset_index())
        weights = subset.loc[subset.peso_grd.notna()].groupby(["year", "definition", "role"]).peso_grd.agg(
            n="size", mean="mean", median="median", q25=lambda s: s.quantile(.25), q75=lambda s: s.quantile(.75)).reset_index()
        add("grd_weight", weights)
        # Co-occurring diagnoses (3-character categories) excluding F84 itself.
        long = subset.melt(id_vars=["year", "definition", "role"], value_vars=DIAGS, var_name="position",
                           value_name="code", ignore_index=False)
        long = long.loc[(long.code != "") & ~long.code.str.startswith("F84")]
        long["code3"] = long.code.str[:3]
        long["code_position"] = np.where(long.position == "DIAGNOSTICO1", "principal", "secundario")
        long = long.reset_index().drop_duplicates(["index", "code3", "code_position"])
        add("codiagnoses", count(long, ["year", "definition", "role", "code_position", "code3"]))
        # Persons per calendar year: first record of the year defines age, sex and comuna.
        persons = subset.loc[subset.valid_id].sort_values(["id", "admission", "discharge"]).drop_duplicates(["year", "id"])
        add("persons_year", count(persons, ["year", "definition", "sex", "age_group", "comuna_norm"], "persons"))
        # Persons per identifier era: first record within the era.
        era_persons = subset.loc[subset.valid_id].sort_values(["id", "admission", "discharge"]).drop_duplicates(["era", "id"])
        add("persons_era", count(era_persons, ["era", "definition", "sex", "age_group", "comuna_norm"], "persons"))
        episodes = subset.loc[subset.valid_id].groupby(["era", "id"]).size().reset_index(name="records")
        episodes["episodes"] = pd.cut(episodes.records, bins=[0, 1, 2, 3, 5, np.inf],
                                      labels=["1", "2", "3", "4-5", "6+"]).astype(str)
        mult = episodes.groupby(["era", "episodes"]).agg(persons=("records", "size"), records=("records", "sum")).reset_index()
        mult.insert(1, "definition", name)
        add("multiplicity", mult)
        # Re-hospitalisation with a code of the definition within 30/90 days of a discharge, same era.
        linked = subset.loc[subset.valid_id & subset.valid_dates].sort_values(["era", "id", "admission", "discharge"]).copy()
        linked["next_admission"] = linked.groupby(["era", "id"]).admission.shift(-1)
        linked["gap"] = (linked.next_admission - linked.discharge).dt.days
        era_end = linked.era.map({era: pd.Timestamp(f"{era_years[-1]}-12-31") for era, era_years in ERAS.items()})
        linked["days_to_end"] = (era_end - linked.discharge).dt.days
        rows = []
        for (era, year), group in linked.groupby(["era", "year"]):
            for horizon in (30, 90):
                eligible = group.loc[group.days_to_end >= horizon]
                rows.append(dict(era=era, definition=name, year=year, horizon_days=horizon, discharges=len(group),
                                 eligible=len(eligible), readmitted=int((eligible.gap <= horizon).sum())))
        add("readmissions", pd.DataFrame(rows))
    # F84 subcodes: a record counts once per distinct subcode and position.
    long = data.melt(id_vars=["year", "sex", "age_group"], value_vars=DIAGS, var_name="position", value_name="code",
                     ignore_index=False)
    long = long.loc[long.code.isin(SUBCODES)]
    long["code_position"] = np.where(long.position == "DIAGNOSTICO1", "principal", "secundario")
    long = long.reset_index().drop_duplicates(["index", "code", "code_position"])
    add("subcodes", count(long, ["year", "code", "code_position", "sex", "age_group"]))
    catalogue = hospital_catalogue(root)
    for name, frames_ in tables.items():
        result = pd.concat(frames_, ignore_index=True)
        if name == "hospitals" and len(catalogue):
            result = result.merge(catalogue, how="left", on="COD_HOSPITAL")
        result.to_csv(output / f"grd_epi_{name}.csv", index=False)
    pd.DataFrame(quality).to_csv(output / "grd_epi_quality.csv", index=False)
    catalogue.to_csv(output / "grd_hospital_catalogue.csv", index=False)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("output_files/consolidacion"))
    parser.add_argument("--years", type=int, nargs="+", default=list(range(2019, 2025)))
    parser.add_argument("--chunksize", type=int, default=100000)
    parser.add_argument("--max-chunks", type=int, default=None, help="Solo para pruebas: limita los bloques leídos por año")
    args = parser.parse_args()
    run(data_root(), args.output, args.years, args.chunksize, args.max_chunks)
